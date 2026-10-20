"""Transcript exporters: .txt, .md, .docx, .vtt, .srt, .json, .csv, .xlsx.

Every writer is a pure projection of ``transcript.json``, so the caption files
and the prose files can never disagree about what was said or when. Re-running
the exporters after a re-timestamp is all it takes to bring every output back
into agreement.

Caption files are built from *word* timings rather than segment timings, which
is what lets a cue break at a sensible place instead of wherever the model
happened to end a segment.
"""

from __future__ import annotations

import copy
import csv
import json
import re
from pathlib import Path

from . import languages, review
from .transcript import (
    DEFAULT_PAUSE_GAP,
    LOW_CONFIDENCE,
    _prob_of,
    full_text,
    segment_text,
    speaker_by_id,
    speaker_roster,
    speakers,
    summary,
    turn_blocks,
    turns,
    word_count,
)

ALL_FORMATS = ("txt", "md", "docx", "vtt", "srt", "json", "csv", "xlsx")


def rendered_turns(doc: dict, settings: dict) -> list:
    """Turns prepared for a writer: label, blocks, and pause markers.

    Every prose writer goes through this, so they cannot disagree about when a
    speaker's name is printed. The rule is one label per turn: because a turn
    is a maximal same-speaker run, a pause in the middle of someone's turn
    breaks the paragraph but does not reprint their name.
    """
    gap = float(settings.get("paragraph_gap", DEFAULT_PAUSE_GAP) or 0.0)
    show_pauses = bool(settings.get("show_pause_markers", True))
    minimum = float(settings.get("pause_marker_seconds", 2.0) or 0.0)
    style = str(settings.get("speaker_label_style", "name") or "name")

    out = []
    for turn in turns(doc, gap=gap):
        speaker = speaker_by_id(doc, turn["speaker_id"])
        if speaker:
            label = speaker["short"] if style == "short" else speaker["name"]
        else:
            label = ""

        blocks = []
        for block in turn_blocks(turn):
            text = " ".join(
                t for t in (segment_text(s) for s in block["segments"]) if t
            ).strip()
            if not text:
                continue
            pause = block["pause_before"]
            blocks.append({
                "text": text,
                "start": float(block["segments"][0].get("start") or 0.0),
                "pause": pause if (show_pauses and pause and pause >= minimum) else None,
            })

        if blocks:
            out.append({
                "label": label,
                "speaker": speaker,
                "start": turn["start"],
                "blocks": blocks,
            })
    return out


def pause_marker(seconds: float) -> str:
    return f"({seconds:.1f}s pause)"


def speaker_legend(doc: dict) -> list:
    """(name, acronym) pairs for the roster, for export headers."""
    return [(s.get("name", ""), s.get("short", "")) for s in speaker_roster(doc)]


def provenance_rows(meta: dict) -> list:
    provenance = meta.get("export_provenance") or {}
    rows = []
    if provenance.get("status"):
        rows.append(("Export status", str(provenance["status"]).replace("_", " ").title()))
    if provenance.get("workflow_status"):
        rows.append((
            "Workflow status",
            str(provenance["workflow_status"]).replace("_", " ").title(),
        ))
    if provenance.get("reviewer"):
        rows.append(("Validated by", str(provenance["reviewer"])))
    if provenance.get("validated_at"):
        rows.append(("Validated at", str(provenance["validated_at"])))
    if provenance.get("exported_at"):
        rows.append(("Exported at", str(provenance["exported_at"])))
    if provenance.get("formats"):
        rows.append(("Export formats", ", ".join(provenance["formats"])))
    return [
        (label, str(value).replace("\r", " ").replace("\n", " "))
        for label, value in rows
    ]

def language_description(info: dict) -> str:
    """Human-readable language line: name, code, and how it was determined.

    A research record should say whether the language was chosen by the person
    running the transcription or inferred by the model, since a wrong
    detection on a short or noisy recording is a known failure mode.
    """
    code = info.get("language") or ""
    if not code:
        return ""
    text = languages.label(code)
    requested = info.get("language_requested") or ""
    if requested == languages.AUTO:
        probability = float(info.get("language_probability") or 0.0)
        text += f", detected automatically ({probability * 100:.0f}% confidence)"
    elif requested:
        text += ", selected by the user"
    status = info.get("output_status") or "original"
    if status == "translated":
        text += "; English machine translation (the text was checked and appears to be English)"
    elif status == "english_source":
        text += "; Translate to English was requested, and the speech is already English"
    elif status == "untranslated":
        text += (
            f"; Translate to English was requested, but the text is still in "
            f"{languages.name(code)}. This is NOT an English translation"
        )
    elif status == "unverified":
        text += "; Translate to English was requested, but English output could not be confirmed"
    return text


def operation_description(info: dict) -> str:
    """What was asked of Whisper, and what the stored text turned out to be."""
    if info.get("operation") == languages.TRANSLATE:
        return "Translate to English (Whisper translate task; a machine translation, not a record of what was said)"
    return "Transcription (text in the language spoken)"


def heading_label(info: dict) -> str:
    """Heading for the body of an export."""
    if (info.get("output_status") or "original") in ("translated", "english_source"):
        return "English translation"
    return "Transcript"


def basename_for(doc: dict) -> str:
    """File stem for a document's exports.

    A confirmed English translation is written as ``translation-en`` so it can
    never be mistaken for, or overwrite, an original-language ``transcript``.
    Text that was requested as a translation but is not confirmed English keeps
    the plain name: the file name must not claim a translation the text does
    not show.
    """
    if languages.output_status(doc) in ("translated", "english_source"):
        return "translation-en"
    return "transcript"


def length_label(info: dict) -> str:
    """Unit for the transcript length figure.

    Scripts written without spaces between words have no word count to report,
    so the figure is a count of characters and must be labelled as one.
    """
    return "Characters" if info.get("spaceless") else "Words"


# ---------------------------------------------------------------------------
# Timestamp formatting
# ---------------------------------------------------------------------------

def vtt_time(seconds: float) -> str:
    """WebVTT uses a dot before milliseconds."""
    return _clock(seconds, ".")


def srt_time(seconds: float) -> str:
    """SubRip uses a comma before milliseconds."""
    return _clock(seconds, ",")


def _clock(seconds: float, sep: str) -> str:
    total_ms = int(round(max(0.0, float(seconds)) * 1000.0))
    hours, rem = divmod(total_ms, 3_600_000)
    minutes, rem = divmod(rem, 60_000)
    secs, ms = divmod(rem, 1000)
    return f"{hours:02d}:{minutes:02d}:{secs:02d}{sep}{ms:03d}"


def stamp(seconds: float) -> str:
    """hh:mm:ss used in readable transcripts."""
    total = int(round(max(0.0, float(seconds))))
    hours, rem = divmod(total, 3600)
    minutes, secs = divmod(rem, 60)
    return f"{hours:02d}:{minutes:02d}:{secs:02d}"


# ---------------------------------------------------------------------------
# Caption cue building
# ---------------------------------------------------------------------------

def build_cues(
    doc: dict,
    max_chars: int = 42,
    max_lines: int = 2,
    max_duration: float = 6.0,
) -> list:
    """Split the transcript into caption cues at word boundaries.

    A cue is closed when adding the next word would exceed the character
    budget (``max_chars`` x ``max_lines``), when the cue would run longer than
    ``max_duration``, or at a segment boundary. Words are never split, so the
    result is always readable.
    """
    max_chars = max(10, int(max_chars))
    max_lines = max(1, int(max_lines))
    spaceless = languages.doc_is_spaceless(doc)
    cues: list = []

    for seg in doc.get("segments", []):
        words = seg.get("words") or []
        speaker = (seg.get("speaker") or "").strip()

        if not words:
            # No word timings (an edited segment awaiting a re-timestamp, or a
            # hand-typed segment): emit the whole segment as one cue rather
            # than dropping it from the captions.
            text = segment_text(seg)
            if text:
                cues.append({
                    "start": float(seg.get("start") or 0.0),
                    "end": float(seg.get("end") or 0.0),
                    "speaker": speaker,
                    "text": text,
                })
            continue

        # A speaker tag is rendered inside the cue, so it must count against
        # the line budget or tagged cues overflow. VTT writes "<v Name>" and
        # SRT writes "Name: "; reserve the longer of the two so one cue list
        # can serve both writers.
        reserve = max(len(f"<v {speaker}>"), len(f"{speaker}: ")) if speaker else 0

        # Where the original text had a space between two words. Only needed
        # for scripts written without word spaces, where it is the sole record
        # of spaces around embedded Latin words and between phrases.
        words = [w for w in words if w.get("w")]
        gaps = (
            languages.word_gaps(segment_text(seg), [w["w"] for w in words])
            if spaceless else None
        )

        current: list = []       # indices into `words`
        for index, word in enumerate(words):
            chosen = current + [index]
            prospective = languages.join_tokens(
                [words[i]["w"] for i in chosen], spaceless,
                [gaps[i] for i in chosen] if gaps else None,
            )
            duration = (
                float(word["end"]) - float(words[current[0]]["start"]) if current else 0.0
            )

            too_long = not fits(prospective, max_chars, max_lines, reserve, spaceless)
            too_slow = duration > float(max_duration)
            if current and (too_long or too_slow):
                cues.append(_cue(words, current, gaps, speaker, spaceless))
                current = []
            current.append(index)

        if current:
            cues.append(_cue(words, current, gaps, speaker, spaceless))

    return cues


def _cue(words: list, indices: list, gaps, speaker: str, spaceless: bool = False) -> dict:
    chosen = [words[i] for i in indices]
    return {
        "start": float(chosen[0]["start"]),
        "end": float(chosen[-1]["end"]),
        "speaker": speaker,
        "text": languages.join_tokens(
            [w["w"] for w in chosen], spaceless,
            [gaps[i] for i in indices] if gaps else None,
        ).strip(),
    }


def wrap_lines(
    text: str,
    max_chars: int,
    max_lines: int | None = None,
    spaceless: bool = False,
) -> list:
    """Greedy wrap at ``max_chars``, breaking only between words.

    ``max_lines`` is accepted for call-site clarity but never used to force an
    overlong line: a single word longer than the limit gets its own line rather
    than being split, because splitting a word mid-character is worse than a
    slightly wide caption. ``build_cues`` is responsible for closing a cue
    before it needs more lines than allowed.

    With ``spaceless`` set (Chinese, Japanese, Thai and similar) there are no
    word spaces to break at, so lines break between characters instead.
    Latin words embedded in the text are still kept whole.
    """
    if spaceless:
        return _wrap_spaceless(text, max_chars)
    lines: list = []
    current = ""
    for word in text.split():
        candidate = f"{current} {word}".strip()
        if current and len(candidate) > max_chars:
            lines.append(current)
            current = word
        else:
            current = candidate
    if current:
        lines.append(current)
    return lines or [""]


def _wrap_spaceless(text: str, max_chars: int) -> list:
    lines: list = []
    current = ""
    for chunk in text.split():
        for index, unit in enumerate(languages.split_units(chunk)):
            separator = " " if (index == 0 and current) else ""
            candidate = current + separator + unit
            if current and len(candidate) > max_chars:
                lines.append(current)
                current = unit
            else:
                current = candidate
    if current:
        lines.append(current)
    return lines or [""]


def fits(
    text: str,
    max_chars: int,
    max_lines: int,
    reserve: int = 0,
    spaceless: bool = False,
) -> bool:
    """True when text wraps into max_lines lines of at most max_chars.

    ``reserve`` accounts for a prefix rendered into the first line, such as a
    speaker tag. Checking the real wrap rather than a characters-times-lines
    budget matters because greedy wrapping is not optimal: 84 characters do not
    always fit into two 42-character lines.
    """
    # The prefix is padded with a non-space filler because wrap_lines splits
    # on whitespace and would otherwise discard leading spaces entirely.
    probe = ("x" * reserve) + text if reserve else text
    return len(wrap_lines(probe, max_chars, spaceless=spaceless)) <= max_lines


# ---------------------------------------------------------------------------
# Writers
# ---------------------------------------------------------------------------

def caption_labels(doc: dict, settings: dict) -> dict:
    """Map a display name to the label captions should carry.

    Captions have very little room, so the acronym is usually preferable to a
    full pseudonym. Returning a mapping keeps build_cues unaware of the roster.
    """
    style = str(settings.get("caption_speaker_style", "short") or "short")
    if style == "none":
        return {s.get("name", ""): "" for s in speaker_roster(doc)}
    if style == "short":
        return {
            s.get("name", ""): (s.get("short") or s.get("name", ""))
            for s in speaker_roster(doc)
        }
    return {}


def write_vtt(doc: dict, dest: Path, settings: dict, meta: dict | None = None) -> Path:
    max_chars = int(settings.get("vtt_max_chars", 42))
    max_lines = int(settings.get("vtt_max_lines", 2))
    max_duration = float(settings.get("vtt_max_duration", 6.0))

    cues = build_cues(doc, max_chars, max_lines, max_duration)
    labels = caption_labels(doc, settings)
    for cue in cues:
        if cue["speaker"] in labels:
            cue["speaker"] = labels[cue["speaker"]]
    out = ["WEBVTT", ""]

    title = (meta or {}).get("title") or ""
    if title:
        # NOTE comments are valid WebVTT and ignored by players.
        out += [f"NOTE {title}", ""]
    provenance = provenance_rows(meta or {})
    if provenance:
        out.append("NOTE Export provenance")
        out.extend(f"{label}: {value}" for label, value in provenance)
        out.append("")

    for index, cue in enumerate(cues, start=1):
        body = cue["text"]
        if cue["speaker"]:
            body = f"<v {cue['speaker']}>{body}"
        lines = wrap_lines(body, max_chars, max_lines, languages.doc_is_spaceless(doc))
        out.append(str(index))
        out.append(f"{vtt_time(cue['start'])} --> {vtt_time(cue['end'])}")
        out.extend(lines)
        out.append("")

    dest.parent.mkdir(parents=True, exist_ok=True)
    dest.write_text("\n".join(out), encoding="utf-8")
    return dest


def write_srt(doc: dict, dest: Path, settings: dict, meta: dict | None = None) -> Path:
    max_chars = int(settings.get("vtt_max_chars", 42))
    max_lines = int(settings.get("vtt_max_lines", 2))
    max_duration = float(settings.get("vtt_max_duration", 6.0))

    cues = build_cues(doc, max_chars, max_lines, max_duration)
    labels = caption_labels(doc, settings)
    for cue in cues:
        if cue["speaker"] in labels:
            cue["speaker"] = labels[cue["speaker"]]
    out = []
    for index, cue in enumerate(cues, start=1):
        body = cue["text"]
        if cue["speaker"]:
            body = f"{cue['speaker']}: {body}"
        out.append(str(index))
        out.append(f"{srt_time(cue['start'])} --> {srt_time(cue['end'])}")
        out.extend(wrap_lines(
            body, max_chars, max_lines, languages.doc_is_spaceless(doc)
        ))
        out.append("")

    dest.parent.mkdir(parents=True, exist_ok=True)
    dest.write_text("\n".join(out), encoding="utf-8")
    return dest


def write_txt(doc: dict, dest: Path, settings: dict, meta: dict | None = None) -> Path:
    include_ts = bool(settings.get("txt_include_timestamps", False))
    include_speakers = bool(settings.get("txt_include_speakers", True))
    gap = float(settings.get("paragraph_gap", 1.5))
    meta = meta or {}

    lines: list = []
    title = meta.get("title") or meta.get("filename") or "Transcript"
    lines.append(title)
    lines.append("=" * len(title))
    lines.append("")

    header = _plain_header(doc, meta)
    if header:
        lines.extend(header)
        lines.append("")

    for turn in rendered_turns(doc, settings):
        for index, block in enumerate(turn["blocks"]):
            if block["pause"]:
                lines.append(pause_marker(block["pause"]))

            prefix_parts = []
            if include_ts:
                prefix_parts.append(f"[{stamp(block['start'])}]")
            # The label goes on the first block only: the speaker has not
            # changed just because they paused.
            if include_speakers and turn["label"] and index == 0:
                prefix_parts.append(f"{turn['label']}:")
            prefix = " ".join(prefix_parts)

            lines.append(f"{prefix} {block['text']}".strip() if prefix else block["text"])
            lines.append("")

    dest.parent.mkdir(parents=True, exist_ok=True)
    dest.write_text("\n".join(lines).rstrip() + "\n", encoding="utf-8")
    return dest


def _plain_header(doc: dict, meta: dict) -> list:
    info = summary(doc)
    validated = (
        (meta.get("export_provenance") or {}).get("workflow_status") == "validated"
    )
    rows = []
    if meta.get("project"):
        rows.append(f"Project: {meta['project']}")
    if meta.get("filename"):
        rows.append(f"Source file: {meta['filename']}")
    rows.append(f"Duration: {stamp(info['duration'])}")
    rows.append(f"Operation: {operation_description(info)}")
    rows.append(f"{length_label(info)}: {info['words']}   Segments: {info['segments']}")
    if info["language"]:
        rows.append(f"Language: {language_description(info)}")
    rows.append(f"Model: {info['model']} (OpenAI Whisper, run locally)")
    legend = speaker_legend(doc)
    if legend:
        rows.append("Speakers: " + ", ".join(
            f"{name} ({short})" if short else name for name, short in legend
        ))
    elif info["speakers"]:
        rows.append(f"Speakers: {', '.join(info['speakers'])}")
    rows.append(
        "Review status: "
        + (
            "validated" if validated else
            "human-reviewed" if info["human_edited"] else
            "machine output, not yet reviewed"
        )
    )
    rows.extend(f"{label}: {value}" for label, value in provenance_rows(meta))
    if meta.get("generated") and not (
        (meta.get("export_provenance") or {}).get("exported_at")
    ):
        rows.append(f"Exported: {meta['generated']}")
    return rows


def write_md(doc: dict, dest: Path, settings: dict, meta: dict | None = None) -> Path:
    gap = float(settings.get("paragraph_gap", 1.5))
    include_ts = bool(settings.get("txt_include_timestamps", False))
    include_speakers = bool(settings.get("txt_include_speakers", True))
    meta = meta or {}
    info = summary(doc)
    validated = (
        (meta.get("export_provenance") or {}).get("workflow_status") == "validated"
    )

    def esc(value) -> str:
        text = str(value if value is not None else "")
        return text.replace('"', '\\"')

    lines = ["---"]
    lines.append(f'title: "{esc(meta.get("title") or meta.get("filename") or "Transcript")}"')
    if meta.get("project"):
        lines.append(f'project: "{esc(meta["project"])}"')
    if meta.get("filename"):
        lines.append(f'source_file: "{esc(meta["filename"])}"')
    lines.append(f"duration_seconds: {info['duration']}")
    lines.append(f"duration: \"{stamp(info['duration'])}\"")
    lines.append(f'operation: "{info["operation"]}"')
    lines.append(f'output_status: "{info["output_status"]}"')
    lines.append(f"words: {info['words']}")
    if info["spaceless"]:
        lines.append('length_unit: "characters"')
    lines.append(f"segments: {info['segments']}")
    if info["language"]:
        lines.append(f'language: "{esc(info["language"])}"')
        lines.append(f'language_name: "{esc(info["language_name"])}"')
        if info["language_requested"] == languages.AUTO:
            lines.append('language_source: "detected"')
            lines.append(f'language_confidence: {info["language_probability"]:.4f}')
        elif info["language_requested"]:
            lines.append('language_source: "selected"')
        if info["output_language"]:
            lines.append(f'output_language: "{esc(info["output_language"])}"')
        if info["output_status"] in ("translated", "english_source"):
            lines.append('translated_to: "en"')
    lines.append(f'model: "{esc(info["model"])}"')
    lines.append("engine: faster-whisper (CTranslate2), executed locally")
    lines.append(
        f"human_reviewed: {str(bool(info['human_edited'] or validated)).lower()}"
    )
    for label, value in provenance_rows(meta):
        key = label.lower().replace(" ", "_")
        lines.append(f'{key}: "{esc(value)}"')
    if info["retimestamped_at"]:
        lines.append(f'retimestamped_at: "{esc(info["retimestamped_at"])}"')
    legend = speaker_legend(doc)
    if legend:
        lines.append("speakers:")
        for name, short in legend:
            lines.append(f'  - name: "{esc(name)}"')
            if short:
                lines.append(f'    short: "{esc(short)}"')
    elif info["speakers"]:
        lines.append("speakers:")
        for name in info["speakers"]:
            lines.append(f'  - name: "{esc(name)}"')
    for key in ("principal_investigator", "irb_protocol", "participant_id"):
        if meta.get(key):
            lines.append(f'{key}: "{esc(meta[key])}"')
    if meta.get("generated"):
        lines.append(f'exported: "{esc(meta["generated"])}"')
    lines.append(
        'processing: "'
        + ("Translated" if info["operation"] == languages.TRANSLATE else "Transcribed")
        + ' offline. No audio or text left this machine."'
    )
    lines.append("---")
    lines.append("")

    lines.append(f"# {meta.get('title') or meta.get('filename') or 'Transcript'}")
    lines.append("")
    if meta.get("description"):
        lines.append(f"_{meta['description']}_")
        lines.append("")

    lines.append(f"## {heading_label(info)}")
    lines.append("")

    for turn in rendered_turns(doc, settings):
        for index, block in enumerate(turn["blocks"]):
            if block["pause"]:
                lines.append(f"_{pause_marker(block['pause'])}_")
                lines.append("")

            parts = []
            if include_speakers and turn["label"] and index == 0:
                parts.append(f"**{turn['label']}**")
            if include_ts:
                parts.append(f"`{stamp(block['start'])}`")
            prefix = " ".join(parts)
            lines.append(f"{prefix} {block['text']}".strip() if prefix else block["text"])
            lines.append("")

    dest.parent.mkdir(parents=True, exist_ok=True)
    dest.write_text("\n".join(lines).rstrip() + "\n", encoding="utf-8")
    return dest


def write_json(doc: dict, dest: Path, settings: dict, meta: dict | None = None) -> Path:
    payload = {
        "meta": meta or {},
        "summary": summary(doc),
        "speakers": speaker_roster(doc),
        "transcript": doc,
    }
    dest.parent.mkdir(parents=True, exist_ok=True)
    dest.write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")
    return dest


def _mark_right_to_left(paragraph) -> None:
    """Make a Word paragraph right-to-left (Arabic, Hebrew and similar).

    Without this Word keeps the letters in the right order but aligns the
    paragraph to the left and lays the speaker label out on the wrong side.
    """
    from docx.oxml import OxmlElement

    paragraph._p.get_or_add_pPr().append(OxmlElement("w:bidi"))
    for run in paragraph.runs:
        run._r.get_or_add_rPr().append(OxmlElement("w:rtl"))


def write_docx(doc: dict, dest: Path, settings: dict, meta: dict | None = None) -> Path:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    gap = float(settings.get("paragraph_gap", 1.5))
    include_ts = bool(settings.get("txt_include_timestamps", False))
    include_speakers = bool(settings.get("txt_include_speakers", True))
    meta = meta or {}
    info = summary(doc)
    validated = (
        (meta.get("export_provenance") or {}).get("workflow_status") == "validated"
    )

    document = Document()

    style = document.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    title = meta.get("title") or meta.get("filename") or "Transcript"
    heading = document.add_heading(title, level=0)
    heading.alignment = WD_ALIGN_PARAGRAPH.LEFT

    subtitle = document.add_paragraph()
    run = subtitle.add_run(
        ("Translated to English" if info["operation"] == languages.TRANSLATE else "Transcribed")
        + " locally with OpenAI Whisper - no data left this machine"
    )
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    # Metadata table: the provenance a research record needs.
    rows = [
        ("Source file", meta.get("filename", "")),
        ("Project", meta.get("project", "")),
        ("Duration", stamp(info["duration"])),
        ("Operation", operation_description(info)),
        (length_label(info), str(info["words"])),
        ("Segments", str(info["segments"])),
        ("Language", language_description(info) or "not recorded"),
        ("Model", info["model"]),
        ("Review status", "Validated" if validated else
         "Human-reviewed" if info["human_edited"]
         else "Machine output, not yet reviewed"),
    ]
    rows.extend(provenance_rows(meta))
    for key in ("principal_investigator", "irb_protocol", "participant_id"):
        if meta.get(key):
            rows.append((key.replace("_", " ").title(), str(meta[key])))
    legend = speaker_legend(doc)
    if legend:
        rows.append(("Speakers", ", ".join(
            f"{name} ({short})" if short else name for name, short in legend
        )))
    elif info["speakers"]:
        rows.append(("Speakers", ", ".join(info["speakers"])))
    if meta.get("generated") and not (
        (meta.get("export_provenance") or {}).get("exported_at")
    ):
        rows.append(("Exported", str(meta["generated"])))

    table = document.add_table(rows=0, cols=2)
    table.style = "Light Grid Accent 1"
    for label, value in rows:
        if not value:
            continue
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = str(value)
        for paragraph in cells[0].paragraphs:
            for r in paragraph.runs:
                r.bold = True
                r.font.size = Pt(9)
        for paragraph in cells[1].paragraphs:
            for r in paragraph.runs:
                r.font.size = Pt(9)

    document.add_paragraph()
    document.add_heading(heading_label(info), level=1)

    rtl = languages.doc_is_rtl(doc)

    for turn in rendered_turns(doc, settings):
        for index, block in enumerate(turn["blocks"]):
            if block["pause"]:
                pause_para = document.add_paragraph()
                pause_run = pause_para.add_run(pause_marker(block["pause"]))
                pause_run.italic = True
                pause_run.font.size = Pt(9)
                pause_run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)

            paragraph = document.add_paragraph()
            if include_ts:
                ts_run = paragraph.add_run(f"[{stamp(block['start'])}] ")
                ts_run.font.size = Pt(9)
                ts_run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
            if include_speakers and turn["label"] and index == 0:
                sp_run = paragraph.add_run(f"{turn['label']}: ")
                sp_run.bold = True
                # Tint the name with the speaker's colour so the document
                # carries the same visual coding as the editor.
                colour = (turn.get("speaker") or {}).get("color")
                if colour and re.fullmatch(r"#[0-9a-fA-F]{6}", colour):
                    sp_run.font.color.rgb = RGBColor.from_string(colour[1:].upper())
            paragraph.add_run(block["text"])
            if rtl:
                _mark_right_to_left(paragraph)

    if meta.get("notes"):
        document.add_heading("Project notes", level=1)
        document.add_paragraph(str(meta["notes"]))

    dest.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(dest))
    return dest


# ---------------------------------------------------------------------------
# Review report (.csv / .xlsx)
# ---------------------------------------------------------------------------
# A spreadsheet with one row per segment, so a team can split up and review a
# transcript together in Excel or Google Sheets instead of each person
# re-running it locally. The machine columns come straight from the transcript
# and the Review Mode queue; the last four columns are left blank for people.
# This is an export only: decisions made in the sheet are not read back.

REVIEW_COLUMNS = (
    "Segment",
    "Start",
    "End",
    "Speaker",
    "Text",
    "Needs review",
    "Review flags",
    "Resolved in Local Scribe",
    "Low-confidence words",
    "Lowest word confidence",
    "Reviewer",
    "Decision",
    "Corrected text",
    "Comments",
)
REVIEWER_COLUMNS = ("Reviewer", "Decision", "Corrected text", "Comments")
REVIEW_DECISIONS = ("Correct as is", "Needs correction", "Unsure / discuss")

# Spreadsheet apps run a cell that starts with one of these as a formula, so
# transcript text such as "=yes" or "-- pause" is written as plain text.
_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


def review_rows(doc: dict) -> list:
    """One dict per segment, in time order, keyed by REVIEW_COLUMNS.

    ``Lowest word confidence`` is a 0-1 float (or None when the segment has
    no timed words); the writers decide how to display it.
    """
    # queue_items() repairs each segment's review state in place, and an
    # export must never change the transcript it describes.
    working = copy.deepcopy(doc)
    by_segment: dict = {}
    for item in review.queue_items(working):
        by_segment.setdefault(item["segment_id"], []).append(item)

    rows = []
    for seg in working.get("segments") or []:
        seg_id = int(seg.get("id", 0))
        items = by_segment.get(seg_id, [])
        open_labels = [i["label"] for i in items if not i["resolved"]]
        done_labels = [i["label"] for i in items if i["resolved"]]

        low_words = []
        lowest = None
        for word in seg.get("words") or []:
            prob = _prob_of(word)
            lowest = prob if lowest is None else min(lowest, prob)
            if prob < LOW_CONFIDENCE:
                token = (word.get("w") or "").strip()
                if not token:
                    continue
                # A human-corrected word is stored with confidence 0.0.
                if prob == 0.0 and seg.get("edited"):
                    low_words.append(f"{token} (corrected)")
                else:
                    low_words.append(f"{token} ({prob * 100:.0f}%)")

        rows.append({
            "Segment": seg_id,
            "Start": stamp(seg.get("start") or 0.0),
            "End": stamp(seg.get("end") or 0.0),
            "Speaker": (seg.get("speaker") or "").strip(),
            "Text": segment_text(seg),
            "Needs review": "Yes" if open_labels else "No",
            "Review flags": "; ".join(open_labels),
            "Resolved in Local Scribe": "; ".join(done_labels),
            "Low-confidence words": ", ".join(low_words),
            "Lowest word confidence": lowest,
            "Reviewer": "",
            "Decision": "",
            "Corrected text": "",
            "Comments": "",
        })
    return rows


def _sheet_text(value) -> str:
    text = "" if value is None else str(value)
    if text.startswith(_FORMULA_PREFIXES):
        return "'" + text
    return text


def write_csv(doc: dict, dest: Path, settings: dict, meta: dict | None = None) -> Path:
    """Review report as CSV.

    Written as UTF-8 with a byte-order mark so Excel opens non-English text
    correctly on a double-click.
    """
    dest.parent.mkdir(parents=True, exist_ok=True)
    with dest.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(REVIEW_COLUMNS)
        for row in review_rows(doc):
            cells = []
            for column in REVIEW_COLUMNS:
                value = row[column]
                if column == "Lowest word confidence":
                    cells.append("" if value is None else f"{value * 100:.0f}%")
                elif column == "Segment":
                    cells.append(str(value))
                else:
                    cells.append(_sheet_text(value))
            writer.writerow(cells)
    return dest


def write_xlsx(doc: dict, dest: Path, settings: dict, meta: dict | None = None) -> Path:
    """Review report as an Excel workbook.

    Sheet "Review" holds the rows, with segments that still need review
    highlighted, a frozen header, filters, and a drop-down for the decision.
    Sheet "About" records what the file was made from and how to use it.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    meta = meta or {}
    rows = review_rows(doc)

    book = Workbook()
    sheet = book.active
    sheet.title = "Review"

    header_font = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="D9E2EC")
    reviewer_fill = PatternFill("solid", fgColor="E6F4EA")
    flag_fill = PatternFill("solid", fgColor="FFF2CC")
    wrap = Alignment(wrap_text=True, vertical="top")

    widths = {
        "Segment": 9, "Start": 10, "End": 10, "Speaker": 14, "Text": 60,
        "Needs review": 9, "Review flags": 22, "Resolved in Local Scribe": 22,
        "Low-confidence words": 30, "Lowest word confidence": 12,
        "Reviewer": 14, "Decision": 18, "Corrected text": 50, "Comments": 40,
    }

    for col, column in enumerate(REVIEW_COLUMNS, start=1):
        cell = sheet.cell(row=1, column=col, value=column)
        cell.font = header_font
        cell.fill = reviewer_fill if column in REVIEWER_COLUMNS else header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        sheet.column_dimensions[get_column_letter(col)].width = widths[column]

    for r, row in enumerate(rows, start=2):
        flagged = row["Needs review"] == "Yes"
        for col, column in enumerate(REVIEW_COLUMNS, start=1):
            value = row[column]
            cell = sheet.cell(row=r, column=col)
            if column == "Segment":
                cell.value = value
                cell.alignment = Alignment(horizontal="center", vertical="top")
            elif column == "Lowest word confidence":
                cell.value = value
                cell.number_format = "0%"
            else:
                cell.value = "" if value is None else str(value)
                # Keep text such as "=yes" from being stored as a formula.
                cell.data_type = "s"
            if column != "Segment":
                cell.alignment = wrap
            if flagged and column not in REVIEWER_COLUMNS:
                cell.fill = flag_fill

    sheet.freeze_panes = "B2"
    last_col = get_column_letter(len(REVIEW_COLUMNS))
    sheet.auto_filter.ref = f"A1:{last_col}{max(1, len(rows) + 1)}"

    if rows:
        decision_col = get_column_letter(REVIEW_COLUMNS.index("Decision") + 1)
        choices = DataValidation(
            type="list",
            formula1='"' + ",".join(REVIEW_DECISIONS) + '"',
            allow_blank=True,
        )
        sheet.add_data_validation(choices)
        choices.add(f"{decision_col}2:{decision_col}{len(rows) + 1}")

    about = book.create_sheet("About")
    lines = [("Local Scribe review report", "")]
    for entry in _plain_header(doc, meta):
        label, _, value = entry.partition(": ")
        lines.append((label, value))
    open_rows = sum(1 for row in rows if row["Needs review"] == "Yes")
    lines += [
        ("Segments needing review", f"{open_rows} of {len(rows)}"),
        ("Low-confidence threshold", f"words below {LOW_CONFIDENCE * 100:.0f}% model confidence"),
        ("", ""),
        ("How to use", "Highlighted rows still have open items in Review Mode. "
                       "Filter 'Needs review' to Yes to see only those."),
        ("", "Fill in Reviewer, Decision, Corrected text and Comments. "
             "Corrections are not imported back automatically: apply them "
             "in the Local Scribe editor."),
        ("Data security", "This file contains the full transcript text. Keep it "
                          "in storage approved for this data."),
    ]
    for r, (label, value) in enumerate(lines, start=1):
        a = about.cell(row=r, column=1, value=label)
        b = about.cell(row=r, column=2, value=value)
        a.data_type = b.data_type = "s"
        a.font = header_font
        a.alignment = b.alignment = wrap
    about["A1"].font = Font(bold=True, size=14)
    about.column_dimensions["A"].width = 26
    about.column_dimensions["B"].width = 90

    dest.parent.mkdir(parents=True, exist_ok=True)
    book.save(str(dest))
    return dest


WRITERS = {
    "txt": write_txt,
    "md": write_md,
    "docx": write_docx,
    "vtt": write_vtt,
    "srt": write_srt,
    "json": write_json,
    "csv": write_csv,
    "xlsx": write_xlsx,
}


def export_all(
    doc: dict,
    out_dir: Path,
    settings: dict,
    meta: dict | None = None,
    formats=None,
    basename: str | None = None,
) -> dict:
    """Write every requested format. Returns {format: path or error string}.

    Unless a ``basename`` is given, the file stem follows the document (see
    :func:`basename_for`) and the other stem's files are removed, so the
    outputs folder always describes the transcript that is current. The
    superseded transcript itself is kept in the document's edit history.
    """
    wanted = list(formats or settings.get("formats") or ALL_FORMATS)
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    automatic = basename is None
    basename = basename or basename_for(doc)

    written: dict = {}
    for fmt in wanted:
        writer = WRITERS.get(fmt)
        if writer is None:
            continue
        dest = out_dir / f"{basename}.{fmt}"
        try:
            written[fmt] = str(writer(doc, dest, settings, meta))
        except Exception as exc:
            written[fmt] = f"error: {exc}"

    if automatic and any(not str(v).startswith("error") for v in written.values()):
        for other in ("transcript", "translation-en"):
            if other != basename:
                for fmt in ALL_FORMATS:
                    (out_dir / f"{other}.{fmt}").unlink(missing_ok=True)
    return written


def preview_stats(doc: dict) -> dict:
    """Small summary used by the UI after an export."""
    info = summary(doc)
    return {
        "words": info["words"],
        "segments": info["segments"],
        "characters": len(full_text(doc)),
        "low_confidence": info["low_confidence"],
        "low_confidence_threshold": LOW_CONFIDENCE,
        "speakers": speakers(doc),
        "word_total": word_count(doc),
    }
