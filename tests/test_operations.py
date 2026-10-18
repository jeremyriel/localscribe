"""Tests for the Transcribe / Translate to English workflow.

Covers, with synthetic data and stub engines only:

* the per-project operation, its default, and its independence from AI Settings
* which catalogue models can translate, and the recommendations per operation
* refusing unsuitable model/operation combinations in the API and in the worker
* terminology on the project page, document page, statuses and exports
* what a requested translation actually produced (English, still the spoken
  language, or unconfirmed) and the proofreading / counting / caption rules
  that follow from it
* existing documents keeping the labels they were made under
* never overwriting a finished result without consent, and keeping both

Nothing here loads a Whisper model or touches real audio. The stub engines
return whatever text a test tells them to, so these tests show that the
application asks for, records and labels results correctly. They do NOT show
that any model translates well; that needs real audio (see the manual tests).

Run with:  .venv/Scripts/python.exe -m tests.test_operations
"""

from __future__ import annotations

import json
import os
import re
import shutil
import sys
import tempfile
from pathlib import Path
from types import SimpleNamespace

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

_TMP = Path(tempfile.mkdtemp(prefix="localscribe-ops-test-")).resolve()
os.environ["LOCALSCRIBE_DATA_DIR"] = str(_TMP)

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except (AttributeError, ValueError):
    pass

from app import catalog, exporters, jobs, languages, projects, proofread  # noqa: E402
from app.config import PATHS, SETTINGS  # noqa: E402
from app.engine import Cancelled  # noqa: E402
from app.transcript import (  # noqa: E402
    make_segment, make_word, new_document, normalise, summary, word_count,
)

assert str(PATHS.projects).startswith(str(_TMP)), "tests must not touch real data"

FAILURES: list = []


def check(label: str, condition: bool, detail: str = "") -> None:
    if condition:
        print(f"  ok   {label}")
    else:
        print(f"  FAIL {label} {ascii(detail)}")
        FAILURES.append(label)


# ---------------------------------------------------------------------------
# Synthetic fixtures
# ---------------------------------------------------------------------------

EN_TOKENS = ["Today", "the", "weather", "is", "nice.", "I", "use", "Zoom", "for", "meetings."]
ES_TOKENS = ["¿Cómo", "está", "usted,", "señora", "Núñez?", "Muy", "bien,", "gracias.",
             "Mañana", "vamos", "al", "niño."]
ZH_TEXT = "今天天气很好。我用 Zoom 开会。"
ZH_TOKENS = ["今天", "天气", "很", "好。", "我", "用", "Zoom", "开会。"]
# Chinese and English mixed: not clearly either, so output cannot be confirmed.
MIXED_TOKENS = ["我用", "Zoom", "开会", "and", "then", "we", "talked", "about", "the",
                "plan", "for", "the", "week", "and", "the", "budget"]

HW_CUDA = {"cpu_cores": 16, "cuda_usable": True, "gpus": ["NVIDIA GeForce RTX 3060"],
           "apple_silicon": False, "mlx_available": False, "unified_memory_gb": None,
           "cpu_compute_types": ["int8"], "warnings": [], "notes": []}
HW_CPU4 = dict(HW_CUDA, cpu_cores=4, cuda_usable=False, gpus=[])
HW_CPU8 = dict(HW_CUDA, cpu_cores=8, cuda_usable=False, gpus=[])
HW_CPU16 = dict(HW_CUDA, cpu_cores=16, cuda_usable=False, gpus=[])


def hw_mlx(gb):
    return dict(HW_CUDA, cuda_usable=False, gpus=[], cpu_cores=10, apple_silicon=True,
                mlx_available=True, unified_memory_gb=gb)


def timed(tokens, start=0.0, step=0.4):
    words, cursor = [], start
    for token in tokens:
        words.append(make_word(token, cursor, cursor + step - 0.05, 0.9))
        cursor += step
    return words


def build_doc(tokens, *, language, operation="", text=None, requested="auto",
              model="medium", legacy_task=None):
    """A stored transcript. ``legacy_task`` mimics a record made before the
    operation was stored: only the Whisper task survives, in the engine block."""
    engine = {"task": legacy_task} if legacy_task else ({"task": operation} if operation else {})
    doc = new_document(duration=8.0, model=model, language=language, language_probability=0.97,
                       language_requested=requested, operation=operation, engine=engine)
    words = timed(tokens)
    doc["segments"] = [make_segment(0, 0.0, words[-1]["end"],
                                    text if text is not None else " ".join(tokens), words)]
    normalise(doc)
    return doc


def original_es():
    return build_doc(ES_TOKENS, language="es", operation="transcribe")


def translated_en():
    return build_doc(EN_TOKENS, language="zh", operation="translate")


def untranslated_es():
    """Translate to English was requested, Spanish came back."""
    return build_doc(ES_TOKENS, language="es", operation="translate")


def untranslated_zh():
    return build_doc(ZH_TOKENS, language="zh", operation="translate", text=ZH_TEXT)


def store(slug, doc, *, status="transcribed", write_exports=True):
    """Put a finished transcript into a project the way the app would."""
    doc_id = projects.add_document(slug, "recording.wav", data=b"RIFF synthetic")["id"]
    projects.save_transcript(slug, doc_id, doc, revision=False)
    projects.set_status(slug, doc_id, status, "done")
    if write_exports:
        exporters.export_all(doc, projects.outputs_dir(slug, doc_id), SETTINGS.all(),
                             {"title": "recording", "project": "p"})
    return doc_id


def new_project(title, **fields):
    return projects.create_project(title, **fields)["slug"]


def outputs(slug, doc_id):
    return sorted(p.name for p in projects.outputs_dir(slug, doc_id).iterdir())


def _client():
    try:
        from fastapi.testclient import TestClient
    except ImportError as exc:
        raise SystemExit(
            "fastapi.testclient needs httpx. Install the test requirements: "
            "pip install -r requirements-dev.txt"
        ) from exc
    from app import update_check
    update_check.refresh_state_async = lambda: None
    from app import main
    return main, TestClient(main.app)


def status_for(model="medium", downloaded=True):
    label = (catalog.model_spec(model).label if catalog.model_spec(model) else model)
    return lambda: {
        "model_downloaded": downloaded, "selected_model": model,
        "selected_model_label": label, "state": "idle", "ready": False,
        "loaded_model": "", "loaded_device": "", "target_device": "cpu",
        "detail": "", "warnings": [], "notes": [], "offline_lock": False,
        "cuda_usable": False, "gpu_count": 0, "gpu_name": "", "driver_version": None,
        "cpu_cores": 8, "backend": None, "apple_silicon": False, "mlx_available": False,
        "unified_memory_gb": None, "target_compute": "int8", "loaded_compute": "",
        "load_seconds": 0,
    }


class Patched:
    """Temporarily replace attributes; restore on exit."""

    def __init__(self, *triples):
        self.triples = triples
        self.saved = []

    def __enter__(self):
        for obj, name, value in self.triples:
            self.saved.append((obj, name, getattr(obj, name)))
            setattr(obj, name, value)
        return self

    def __exit__(self, *exc):
        for obj, name, value in reversed(self.saved):
            setattr(obj, name, value)


# ---------------------------------------------------------------------------
# A stub engine and a way to run one job through the real pipeline
# ---------------------------------------------------------------------------

class StubEngine:
    """Stands in for Whisper inside the job runner. It returns ``texts``
    whatever it is asked to do, as a model that cannot translate would."""

    def __init__(self, texts, detected="es", raises=None):
        self.texts = texts
        self.detected = detected
        self.raises = raises
        self.overrides = None

    def transcribe(self, pcm, *, duration, overrides=None, on_segment=None, should_cancel=None):
        self.overrides = dict(overrides or {})
        if self.raises:
            raise self.raises
        tracker = SimpleNamespace(snapshot=lambda: {
            "progress": 0.5, "eta_text": "1s", "words_per_second": 2.0,
            "tokens_per_second": 4.0, "realtime_factor": 3.0})
        segments = []
        for index, text in enumerate(self.texts):
            words = [SimpleNamespace(word=f" {t}", start=index * 3 + i * 0.4,
                                     end=index * 3 + i * 0.4 + 0.3, probability=0.9)
                     for i, t in enumerate(text.split())]
            seg = SimpleNamespace(start=index * 3.0, end=index * 3.0 + 2.5, text=f" {text}",
                                  words=words, tokens=[1, 2], avg_logprob=-0.2,
                                  no_speech_prob=0.01, compression_ratio=1.1, temperature=0.0)
            segments.append(seg)
            on_segment(index, seg, tracker)
        forced = self.overrides.get("language")
        spoken = forced if forced and forced != "auto" else self.detected
        return {
            "segments": segments, "info": SimpleNamespace(language=spoken),
            "language": spoken,
            "language_probability": 1.0 if forced and forced != "auto" else 0.97,
            "language_requested": forced if forced and forced != "auto" else "auto",
            "stats": {"wall_seconds": 1.0, "realtime_factor": 3.0, "model": SETTINGS.get("model"),
                      "device": "cpu", "compute_type": "int8"},
            "params": {"task": self.overrides.get("task", "transcribe")},
        }


def run_job(slug, doc_id, *, operation, texts, language="auto", detected="es", raises=None):
    stub = StubEngine(texts, detected=detected, raises=raises)
    info = SimpleNamespace(
        container="wav", audio_codec="pcm", sample_rate=16000, channels=1, duration=6.0,
        kind="audio", browser_playable=True, has_video=False, video_codec=None,
        width=None, height=None, to_dict=lambda: {"container": "wav", "has_video": False})
    queue = jobs.JobQueue()
    with Patched(
        (jobs, "ENGINE", stub),
        (jobs.media, "probe", lambda *a, **k: info),
        (jobs.media, "decode_to_pcm", lambda *a, **k: [0.0] * (16000 * 6)),
        (jobs.media, "write_waveform_cache", lambda *a, **k: None),
        (jobs.media, "describe", lambda i: "synthetic"),
    ):
        job = jobs.Job("transcribe", "recording.wav", slug=slug, doc_id=doc_id,
                       language=language, operation=operation)
        try:
            queue._do_transcription(job)
        except Cancelled:
            raise
        except Exception as exc:                       # what the worker loop does
            projects.set_status(slug, doc_id, "failed", str(exc), error=str(exc))
            return stub, str(exc)
    return stub, None


EN_JOB_TEXTS = ["Today the weather is nice and we are in the office.",
                "I use Zoom for meetings with the team."]
ES_JOB_TEXTS = ["¿Cómo está usted, señora Núñez?", "Muy bien, gracias. Mañana vamos al niño."]


# ---------------------------------------------------------------------------
# What a requested translation produced
# ---------------------------------------------------------------------------

def test_english_output_check():
    check("English prose is English", languages.assess_english(
        "Today the weather is nice, and I think we should go to the park with them.") == "english")
    check("conversational English with fillers is English", languages.assess_english(
        "Yeah, so I was just saying that you know it was not what they wanted at all.") == "english")
    check("Spanish with accents is not English", languages.assess_english(
        "¿Cómo está usted, señora Núñez? Muy bien, gracias. Mañana vamos al niño.") == "not_english")
    check("French is not English", languages.assess_english(
        "Bonjour, comment allez-vous aujourd'hui? Je voudrais un café, s'il vous plaît.") == "not_english")
    check("German is not English", languages.assess_english(
        "Das ist ein Test und wir sprechen heute über das Wetter in der Stadt.") == "not_english")
    check("Chinese is not English", languages.assess_english(ZH_TEXT) == "not_english")
    check("Arabic is not English", languages.assess_english(
        "مرحبا بكم في المقابلة اليوم سوف نتحدث عن التجربة") == "not_english")
    check("a very short text cannot be judged", languages.assess_english("Okay then.") == "unknown")
    check("mixed Chinese and English is only uncertain", languages.assess_english(
        "我用 Zoom 开会 and then we talked about the plan for the week and the budget") == "uncertain")
    check("text that is neither clearly English nor not is uncertain, not a guess",
          languages.assess_english("Dit is een test en we praten vandaag over het weer in de stad.")
          in ("uncertain", "not_english"))


def test_output_status_and_language():
    cases = [
        ("a transcription", original_es(), "original", "es"),
        ("a translation that is English", translated_en(), "translated", "en"),
        ("a translation that is still Spanish", untranslated_es(), "untranslated", "es"),
        ("a translation that is still Chinese", untranslated_zh(), "untranslated", "zh"),
        ("English speech translated", build_doc(EN_TOKENS, language="en", operation="translate"),
         "english_source", "en"),
    ]
    for label, doc, status, language in cases:
        check(f"{label}: status {status}", languages.output_status(doc) == status,
              languages.output_status(doc))
        check(f"{label}: text language {language}", languages.output_language(doc) == language,
              languages.output_language(doc))

    mixed = build_doc(["我用", "Zoom", "开会", "and", "then", "we", "talked", "about", "the",
                       "plan", "for", "the", "week", "and", "the", "budget"],
                      language="zh", operation="translate")
    check("unconfirmed output is unverified", languages.output_status(mixed) == "unverified")
    check("and its language is undetermined, not English",
          languages.output_language(mixed) == languages.UNDETERMINED)
    check("the request alone proves nothing",
          languages.output_language(build_doc(ES_TOKENS, language="es", operation="translate")) != "en")


def test_text_rules_follow_the_actual_text():
    es, zh, en = untranslated_es(), untranslated_zh(), translated_en()
    check("untranslated Spanish is not proofread as English", not proofread.supports_document(es))
    check("untranslated Chinese is not proofread as English", not proofread.supports_document(zh))
    check("a real English translation is proofread", proofread.supports_document(en))

    check("untranslated Chinese is counted in characters, not as one word",
          word_count(zh) == 14 and summary(zh)["spaceless"] is True, str(word_count(zh)))
    check("a real English translation is counted in words", word_count(en) == 10)
    check("untranslated Spanish is counted in words", word_count(es) == 12)

    cues = exporters.build_cues(zh)
    check("untranslated Chinese captions keep Chinese spacing",
          cues[0]["text"] == ZH_TEXT, cues[0]["text"])
    cues = exporters.build_cues(en)
    check("English translation captions use spaces",
          cues[0]["text"] == " ".join(EN_TOKENS), cues[0]["text"])

    mixed = build_doc(MIXED_TOKENS, language="zh", operation="translate")
    check("undetermined text is not proofread", not proofread.supports_document(mixed))
    check("and is described honestly", summary(mixed)["output_language_name"] == "Undetermined")

    arabic = build_doc(["مرحبا", "بكم", "في", "المقابلة؟"], language="ar", operation="translate",
                       text="مرحبا بكم في المقابلة اليوم سوف نتحدث عن التجربة")
    check("untranslated Arabic keeps right-to-left handling", summary(arabic)["rtl"] is True)


# ---------------------------------------------------------------------------
# Which models can translate, and what to recommend
# ---------------------------------------------------------------------------

def test_model_classification():
    unsupported = {
        "tiny.en", "base.en", "small.en", "distil-small.en", "medium.en", "distil-large-v3",
        "large-v3-turbo", "tiny.en-mlx", "base.en-mlx", "small.en-mlx", "medium.en-mlx",
        "large-v3-turbo-mlx", "large-v3-turbo-mlx-4bit",
    }
    limited = {"tiny", "base", "tiny-mlx", "base-mlx"}
    ok = {"small", "medium", "large-v2", "large-v3", "small-mlx", "medium-mlx",
          "medium-mlx-4bit", "large-v3-mlx", "large-v3-mlx-4bit"}
    every = {m.key for m in catalog.CATALOG}
    check("every catalogue model is classified on purpose (a new one forces a decision)",
          unsupported | limited | ok == every and not (unsupported & limited) and not (limited & ok),
          str(every - (unsupported | limited | ok)))

    got = {m.key: catalog.translation_support(m)["level"] for m in catalog.CATALOG}
    check("unsupported models are exactly Turbo and English-only",
          {k for k, v in got.items() if v == "unsupported"} == unsupported)
    check("Tiny and Base are allowed but flagged as poor",
          {k for k, v in got.items() if v == "limited"} == limited)
    check("standard multilingual models are fine",
          {k for k, v in got.items() if v == "ok"} == ok)
    check("every unsupported model explains why",
          all(catalog.translation_support(catalog.model_spec(k))["reason"] for k in unsupported))

    reason = catalog.translation_support(catalog.model_spec("large-v3-turbo"))["reason"]
    check("Turbo is explained as not trained for translation", "not trained for translation" in reason)
    reason = catalog.translation_support(catalog.model_spec("medium.en"))["reason"]
    check("English-only models are explained", "only understands English" in reason)


def test_recommendations():
    cases = [
        ("NVIDIA GPU", HW_CUDA, "large-v3-turbo", "medium"),
        ("CPU, 4 cores", HW_CPU4, "base", "small"),
        ("CPU, 8 cores", HW_CPU8, "small", "small"),
        ("CPU, 16 cores", HW_CPU16, "medium", "medium"),
        ("Apple 64 GB", hw_mlx(64), "large-v3-mlx", "large-v3-mlx"),
        ("Apple 24 GB", hw_mlx(24), "large-v3-turbo-mlx", "medium-mlx"),
        ("Apple 8 GB", hw_mlx(8), "small-mlx", "small-mlx"),
    ]
    for label, hw, transcribe_model, translate_model in cases:
        plain = catalog.recommend(hw)
        translate = catalog.recommend(hw, task="translate")
        check(f"{label}: transcription recommendation is unchanged ({transcribe_model})",
              plain["model"] == transcribe_model, plain["model"])
        check(f"{label}: translation recommendation is {translate_model}",
              translate["model"] == translate_model, translate["model"])
        spec = catalog.model_spec(translate["model"])
        check(f"{label}: the translation recommendation can translate",
              catalog.translation_support(spec)["level"] == "ok")
        check(f"{label}: device and precision are kept",
              (translate["device"], translate["compute_type"]) == (plain["device"], plain["compute_type"]))
        check(f"{label}: the reason says why not Turbo", "not recommended" in translate["reason"]
              or "not trained for translation" in translate["reason"])
        check(f"{label}: alternatives can all translate",
              all(catalog.translation_support(catalog.model_spec(a["model"]))["level"] != "unsupported"
                  for a in translate["alternatives"]))

    payload = catalog.catalog_payload(HW_CPU8)
    flagged = [m["key"] for m in payload["models"] if m["recommended_translate"]]
    check("exactly one model is recommended for translation", flagged == ["small"], str(flagged))
    check("the translation entry reports why Turbo is unsuitable",
          next(m for m in payload["models"] if m["key"] == "large-v3-turbo")["translation"]["level"]
          == "unsupported")
    check("the transcription recommendation is still reported separately",
          payload["recommendation"]["model"] == "small"
          and payload["recommendation_translate"]["model"] == "small")

    with Patched((catalog, "is_downloaded", lambda key: key == "medium")):
        suggestion = catalog.suggest_translation_model(HW_CUDA)
    check("a suggestion names the model and whether it is downloaded",
          suggestion["key"] == "medium" and suggestion["downloaded"] is True and suggestion["label"] == "Medium")
    with Patched((catalog, "is_downloaded", lambda key: False)):
        problem = catalog.operation_problem("large-v3-turbo", "translate", HW_CPU8)
    check("the problem message carries the suggestion without selecting or downloading anything",
          "not trained for translation" in problem and "Suggested: Small (not downloaded yet)" in problem,
          problem)
    check("transcription is never blocked by the model",
          catalog.operation_problem("large-v3-turbo", "transcribe", HW_CPU8) is None
          and catalog.operation_problem("tiny.en", "transcribe", HW_CPU8) is None)
    check("limited models are allowed to run (with a warning, not a block)",
          catalog.operation_problem("tiny", "translate", HW_CPU8) is None)


# ---------------------------------------------------------------------------
# The operation is per project
# ---------------------------------------------------------------------------

def test_operation_default_and_persistence():
    SETTINGS.update({"task": "transcribe", "language": "auto", "model": "medium"})
    a = new_project("Operacion A")
    check("a new project saves the AI Settings default explicitly",
          projects.load_project(a)["operation"] == "transcribe")

    SETTINGS.update({"task": "translate"})
    b = new_project("Operacion B")
    check("a new project starts from a translate default",
          projects.load_project(b)["operation"] == "translate")
    check("but an existing project does not follow the default",
          projects.project_operation(projects.load_project(a)) == "transcribe")

    path = projects.project_dir(a) / projects.PROJECT_FILE
    legacy = projects.read_json(path)
    legacy.pop("operation", None)
    projects.write_json(path, legacy)
    check("a project saved before operations existed is Transcribe, whatever the default",
          projects.project_operation(projects.load_project(a)) == "transcribe")
    check("and with no saved value the record says so honestly",
          "operation" not in projects.load_project(a))
    SETTINGS.update({"task": "transcribe"})

    projects.save_project(a, {"operation": "Translate to English"})
    check("the project's choice is saved (names are understood)",
          projects.load_project(a)["operation"] == "translate")
    check("the AI Settings default is untouched by it", SETTINGS.get("task") == "transcribe")
    projects.save_project(a, {"operation": "nonsense"})
    check("an invalid value is ignored", projects.load_project(a)["operation"] == "translate")
    projects.save_project(a, {"title": "Operacion A bis"})
    check("saving other details keeps the operation", projects.load_project(a)["operation"] == "translate")
    check("another project is unaffected", projects.load_project(b)["operation"] == "translate"
          and projects.project_operation(projects.load_project(new_project("Operacion C"))) == "transcribe")

    check("the default setting is a dropdown with readable labels",
          any(f["key"] == "task" and f["option_labels"]["translate"] == "Translate to English"
              for g in __import__("app.config", fromlist=["SETTINGS_SCHEMA"]).SETTINGS_SCHEMA
              for f in g["fields"]))


def test_switching_operation_does_not_touch_defaults_or_documents():
    main, client = _client()
    SETTINGS.update({"task": "transcribe", "model": "medium"})
    slug = new_project("Cambio de modo")
    doc_id = store(slug, original_es())
    before = projects.load_document(slug, doc_id)
    transcript_before = (projects.transcript_path(slug, doc_id)).read_bytes()
    files_before = outputs(slug, doc_id)

    r = client.put(f"/api/projects/{slug}", json={"operation": "translate"})
    check("the project switches to translate", r.json()["project"]["operation"] == "translate")
    check("AI Settings default is unchanged", SETTINGS.get("task") == "transcribe")

    after = projects.load_document(slug, doc_id)
    check("the document record is not relabelled",
          after["operation"] == "transcribe" and after["output_status"] == "original"
          and after["language_name"] == before["language_name"])
    check("the transcript file is byte-for-byte unchanged",
          (projects.transcript_path(slug, doc_id)).read_bytes() == transcript_before)
    check("its exports are untouched", outputs(slug, doc_id) == files_before)

    docs = projects.list_documents(slug)
    check("the project's recording list still says original", docs[0]["output_status"] == "original")


# ---------------------------------------------------------------------------
# The API refuses unsuitable combinations
# ---------------------------------------------------------------------------

def test_api_enforces_model_and_operation():
    main, client = _client()
    SETTINGS.update({"language": "auto"})
    slug = new_project("Modelos")
    doc_id = projects.add_document(slug, "a.wav", data=b"RIFF synthetic")["id"]
    url = f"/api/projects/{slug}/transcribe"

    def attempt(model, **body):
        projects.set_status(slug, doc_id, "uploaded", "")
        with Patched((main.ENGINE, "status", status_for(model)),
                     (main.ENGINE, "hardware", lambda refresh=False: HW_CPU8)):
            return client.post(url, json={"documents": [doc_id], **body})

    pending_before = main.QUEUE.pending_count()
    r = attempt("large-v3-turbo", operation="translate")
    check("Turbo + Translate to English is refused", r.status_code == 409, r.text)
    check("the refusal says why and suggests a model",
          "not trained for translation" in r.json()["detail"] and "Suggested:" in r.json()["detail"],
          r.text)
    check("nothing entered the queue", main.QUEUE.pending_count() == pending_before)
    check("and the recording was not marked as queued",
          projects.load_document(slug, doc_id)["status"] == "uploaded")

    r = attempt("medium.en", operation="translate")
    check("an English-only model cannot translate", r.status_code == 409
          and "only understands English" in r.json()["detail"], r.text)
    r = attempt("large-v3-turbo-mlx", operation="translate")
    check("Turbo on Apple GPU is refused too", r.status_code == 409, r.text)

    r = attempt("large-v3-turbo", operation="transcribe")
    check("the same model still transcribes", r.status_code == 200 and len(r.json()["queued"]) == 1, r.text)
    job = main.QUEUE.get(r.json()["queued"][0]["job"])
    check("and the job records the operation", job.payload["operation"] == "transcribe")

    r = attempt("medium", operation="translate")
    check("a translation-capable model is accepted", r.status_code == 200 and len(r.json()["queued"]) == 1, r.text)
    job = main.QUEUE.get(r.json()["queued"][0]["job"])
    check("the job carries translate", job.payload["operation"] == "translate")
    check("the recording is marked as pending a translation",
          projects.load_document(slug, doc_id)["pending_operation"] == "translate")
    check("the response says which operation was queued", r.json()["operation"] == "translate")

    r = attempt("tiny", operation="translate")
    check("a weak model is allowed (it is flagged in the page, not blocked)", r.status_code == 200, r.text)


def test_api_operation_resolution_and_errors():
    main, client = _client()
    SETTINGS.update({"task": "transcribe", "language": "auto"})
    slug = new_project("Resolucion")
    projects.save_project(slug, {"operation": "translate"})
    doc_id = projects.add_document(slug, "a.wav", data=b"RIFF synthetic")["id"]
    url = f"/api/projects/{slug}/transcribe"

    def attempt(**body):
        projects.set_status(slug, doc_id, "uploaded", "")
        with Patched((main.ENGINE, "status", status_for("medium")),
                     (main.ENGINE, "hardware", lambda refresh=False: HW_CPU8)):
            return client.post(url, json={"documents": [doc_id], **body})

    r = attempt()
    check("with nothing sent the project's operation is used",
          main.QUEUE.get(r.json()["queued"][0]["job"]).payload["operation"] == "translate")
    r = attempt(operation="transcribe")
    check("an explicit operation beats the project's",
          main.QUEUE.get(r.json()["queued"][0]["job"]).payload["operation"] == "transcribe")
    r = attempt(operation="Translate")
    check("operation names are normalised",
          main.QUEUE.get(r.json()["queued"][0]["job"]).payload["operation"] == "translate")

    r = attempt(operation="summarise")
    check("an unknown operation is a 400", r.status_code == 400 and "not a supported operation" in r.json()["detail"], r.text)
    r = attempt(language="klingon")
    check("an unknown language is still a 400", r.status_code == 400)
    with Patched((main.ENGINE, "status", status_for("medium", downloaded=False))):
        r = client.post(url, json={"operation": "translate"})
    check("a missing model is still reported first", r.status_code == 409 and "not been downloaded" in r.json()["detail"])
    r = client.post("/api/projects/does-not-exist/transcribe", json={"operation": "translate"})
    check("an unknown project is a 404", r.status_code == 404)
    r = client.put(f"/api/projects/{slug}", json={"operation": "summarise"})
    check("saving an invalid operation does not error and changes nothing",
          r.status_code == 200 and projects.load_project(slug)["operation"] == "translate")
    r = client.post(url, content="not json", headers={"content-type": "application/json"})
    check("a malformed body is handled", r.status_code in (200, 409), r.text)


def test_worker_refuses_a_model_that_changed_while_queued():
    SETTINGS.update({"model": "large-v3-turbo"})
    slug = new_project("Guardia")
    doc_id = projects.add_document(slug, "a.wav", data=b"RIFF synthetic")["id"]
    stub, error = run_job(slug, doc_id, operation="translate", texts=EN_JOB_TEXTS)
    SETTINGS.update({"model": "medium"})
    check("the worker refuses Turbo for a translation", error is not None and "not trained for translation" in error, str(error))
    check("without running the engine", stub.overrides is None)
    check("the recording shows the failure", projects.load_document(slug, doc_id)["status"] == "failed")
    check("and no transcript was written", not projects.transcript_path(slug, doc_id).exists())


def test_global_default_does_not_leak_into_a_run():
    SETTINGS.update({"task": "translate", "model": "medium"})
    slug = new_project("Sin fuga")
    projects.save_project(slug, {"operation": "transcribe"})
    doc_id = projects.add_document(slug, "a.wav", data=b"RIFF synthetic")["id"]
    stub, error = run_job(slug, doc_id, operation="transcribe", texts=ES_JOB_TEXTS)
    SETTINGS.update({"task": "transcribe"})
    check("a transcribe run sends task=transcribe even if the default is translate",
          error is None and stub.overrides.get("task") == "transcribe", str(stub.overrides))
    doc = projects.load_transcript(slug, doc_id)
    check("and records it", doc["operation"] == "transcribe" and doc["engine"]["task"] == "transcribe")


# ---------------------------------------------------------------------------
# What the job records and says
# ---------------------------------------------------------------------------

def test_job_records_what_happened():
    SETTINGS.update({"model": "medium"})
    slug = new_project("Registro")

    doc_id = projects.add_document(slug, "a.wav", data=b"RIFF synthetic")["id"]
    stub, error = run_job(slug, doc_id, operation="translate", texts=EN_JOB_TEXTS, detected="zh")
    doc = projects.load_transcript(slug, doc_id)
    record = projects.load_document(slug, doc_id)
    check("translate reaches the engine as the task", stub.overrides.get("task") == "translate")
    check("English output is recorded as a translation",
          doc["operation"] == "translate" and record["output_status"] == "translated"
          and record["output_language"] == "en", str(record))
    check("the spoken language and how it was found are kept",
          doc["language"] == "zh" and doc["language_requested"] == "auto")
    check("the model actually used is kept", doc["model"] == "medium")
    check("what the text looked like is kept", doc["output_check"]["verdict"] == "english"
          and doc["output_check"]["status"] == "translated" and doc["output_check"]["checked"])
    check("the status line says English translation",
          "English translation ready for review" in record["status_detail"], record["status_detail"])
    check("exports are named as a translation",
          outputs(slug, doc_id) == sorted(f"translation-en.{e}" for e in ("docx", "json", "md", "srt", "txt", "vtt")),
          str(outputs(slug, doc_id)))

    doc_id = projects.add_document(slug, "b.wav", data=b"RIFF synthetic")["id"]
    stub, error = run_job(slug, doc_id, operation="translate", texts=ES_JOB_TEXTS, detected="es")
    doc = projects.load_transcript(slug, doc_id)
    record = projects.load_document(slug, doc_id)
    check("Spanish returned for a translation is recorded as such",
          record["output_status"] == "untranslated" and record["output_language"] == "es"
          and doc["output_check"]["verdict"] == "not_english", str(record))
    check("the status line does not claim a translation",
          "Not translated" in record["status_detail"] and "es" not in record["status_detail"].lower().replace("spanish", ""),
          record["status_detail"])
    check("its exports keep the plain transcript name",
          all(n.startswith("transcript.") for n in outputs(slug, doc_id)))

    doc_id = projects.add_document(slug, "c.wav", data=b"RIFF synthetic")["id"]
    run_job(slug, doc_id, operation="transcribe", texts=ES_JOB_TEXTS, detected="es")
    record = projects.load_document(slug, doc_id)
    check("a transcription is recorded as one",
          record["operation"] == "transcribe" and record["output_status"] == "original"
          and "ready for review" in record["status_detail"] and "translat" not in record["status_detail"].lower())
    check("with an empty output check", projects.load_transcript(slug, doc_id)["output_check"] == {})

    heartbeat = jobs.Job("transcribe", "x", slug="s", doc_id="d", operation="translate").to_dict()
    check("the queue reports the operation of a job", heartbeat["operation"] == "translate")


# ---------------------------------------------------------------------------
# Terminology on every surface
# ---------------------------------------------------------------------------

def test_terms_table():
    t, tr = languages.terms("transcribe"), languages.terms("translate")
    check("transcribe terms", (t["action"], t["progressive"], t["done"], t["result"]) ==
          ("Transcribe", "Transcribing", "Transcribed", "transcript"))
    check("translate terms", (tr["action"], tr["progressive"], tr["done"], tr["result"]) ==
          ("Translate to English", "Translating", "Translated", "English translation"))
    check("both define every key", set(t) == set(tr))
    check("unknown falls back to transcribe", languages.terms("nonsense") == t and languages.terms(None) == t)
    check("the translate explanation says it is a machine translation and not Turbo",
          "machine translation" in tr["explain"] and "not Turbo" in tr["explain"])


def label_of(html, element_id):
    match = re.search(rf'id="{element_id}"[^>]*>(.*?)</button>', html, re.S)
    text = re.sub(r"<[^>]+>", " ", match.group(1)) if match else ""
    return " ".join(text.split())


def test_project_page_follows_the_operation():
    main, client = _client()
    SETTINGS.update({"task": "transcribe", "model": "large-v3-turbo", "language": "auto"})
    slug = new_project("Pagina")
    doc_id = store(slug, original_es())
    with Patched((main.ENGINE, "status", status_for("large-v3-turbo")),
                 (main.ENGINE, "hardware", lambda refresh=False: HW_CPU8)):
        page = client.get(f"/project/{slug}").text
        check("a transcribe project shows Transcribe", label_of(page, "transcribe-btn").startswith("Transcribe")
              and "Translate to English" not in label_of(page, "transcribe-btn"), label_of(page, "transcribe-btn"))
        check("its Transcribe choice is selected", re.search(r'value="transcribe"\s+checked', page) is not None)
        check("its summary explains transcription", "Writes what was said, in the language it was said" in page)
        check("its empty state talks about Transcribe", "then press Transcribe." in page)

        projects.save_project(slug, {"operation": "translate"})
        page = client.get(f"/project/{slug}").text
        check("a translate project shows Translate to English",
              label_of(page, "transcribe-btn").startswith("Translate to English"), label_of(page, "transcribe-btn"))
        check("its Translate choice is selected", re.search(r'value="translate"\s+checked', page) is not None)
        check("its summary says machine translation to English",
              "machine translation" in page and "You will get: English text (machine translation)" in page)
        check("its summary shows the spoken language and the model",
              "Spoken language: Detect automatically" in page and "Model: Large v3 Turbo" in page)
        check("it says the model is unsuitable", "(not suitable for this)" in page)
        check("its empty state talks about Translate to English", "then press Translate to English." in page)
        check("both operations are offered as a labelled group",
              '<legend>What do you want to do?</legend>' in page and page.count('name="operation"') == 2)

        info = main.operation_info()
        check("the page hands the browser the compatibility for both operations",
              info["compat"]["translate"]["ok"] is False and info["compat"]["transcribe"]["ok"] is True)
        check("the translate explanation names the problem and a way forward",
              "not trained for translation" in info["compat"]["translate"]["message"]
              and info["suggestion"]["key"] == "small")
        check("transcription has no warning for Turbo", info["compat"]["transcribe"]["message"] == "")
        check("the browser also gets every term it needs", set(info["terms"]) == {"transcribe", "translate"})
        check("the page embeds that compatibility for the live switch",
              "not trained for translation" in page and '"suggestion"' in page)

    with Patched((main.ENGINE, "status", status_for("medium")),
                 (main.ENGINE, "hardware", lambda refresh=False: HW_CPU8)):
        info = main.operation_info()
        page = client.get(f"/project/{slug}").text
    check("a suitable model gives no warning", info["compat"]["translate"] == {"ok": True, "level": "ok", "message": ""})
    check("and the summary does not call it unsuitable", "(not suitable for this)" not in page)

    with Patched((main.ENGINE, "status", status_for("tiny")),
                 (main.ENGINE, "hardware", lambda refresh=False: HW_CPU8)):
        info = main.operation_info()
    check("a weak model is flagged but allowed",
          info["compat"]["translate"]["ok"] is True and info["compat"]["translate"]["level"] == "limited"
          and "too error-prone" in info["compat"]["translate"]["message"])

    # The project list keeps the labels the documents were made under.
    check("the recording's stored label is still Transcript",
          projects.list_documents(slug)[0]["output_status"] == "original")
    check("and the page did not change that", doc_id == projects.list_documents(slug)[0]["id"])


def test_document_page_labels_follow_the_document():
    main, client = _client()
    SETTINGS.update({"task": "transcribe", "model": "medium"})
    slug = new_project("Documentos")
    projects.save_project(slug, {"operation": "translate"})     # the project's mode is not the document's

    def page(doc):
        doc_id = store(slug, doc)
        return client.get(f"/project/{slug}/document/{doc_id}").text

    original = page(original_es())
    check("an original transcript is headed Transcript", "<h2 style=\"margin:0\">Transcript</h2>" in original)
    check("its downloads say transcripts", "<h3>Download transcripts</h3>" in original)
    check("it carries no translation notice", 'id="output-notice"' not in original)
    check("its facts say Transcribe", "Transcribe<br>" in " ".join(original.split()).replace("Transcribe <br>", "Transcribe<br>"))
    check("even though the project is set to translate", "Translate to English</dd>" not in original)

    translated = page(translated_en())
    check("an English translation is headed English translation",
          "<h2 style=\"margin:0\">English translation</h2>" in translated)
    check("its downloads say English translation", "<h3>Download English translation</h3>" in translated)
    check("it explains it is a machine translation", "English machine translation" in translated)
    check("its find box says translation", 'placeholder="Find in this translation"' in translated)
    check("its facts say what was asked and what came back",
          "Translate to English" in translated and "English translation (machine)" in translated)
    check("its facts show the spoken and the text language",
          "Chinese (zh)" in translated and "English (en)" in translated)
    check("it hands the editor the right noun", 'noun: "translation"' in translated)
    flat = " ".join(translated.split())
    check("its Status fact describes the operation, not the internal state",
          '<dt>Status</dt> <dd id="fact-status">English translation (machine)</dd>' in flat
          and "transcribed</dd>" not in flat)
    check("an original transcript's Status still says Transcribed",
          '<dd id="fact-status">Transcribed</dd>' in " ".join(original.split()))
    check("translated text shows no word-confidence guidance",
          "showConfidence: false" in translated and 'id="jump-low"' not in translated
          and "Uncertain words" not in translated and "low model confidence" not in translated)
    check("it tells the reader why", "Word confidence is not shown" in flat)
    check("a transcript keeps its confidence guidance",
          "showConfidence: true" in original and 'id="jump-low"' in original
          and "Uncertain words" in original)
    low_words = translated_en()
    low_words["segments"][0]["words"][0]["prob"] = 0.1
    check("low-probability words are not counted for a translation",
          summary(low_words)["low_confidence"] == 0 and not summary(low_words)["confidence_meaningful"])
    low_original = original_es()
    low_original["segments"][0]["words"][0]["prob"] = 0.1
    check("but are counted for a transcript",
          summary(low_original)["low_confidence"] == 1 and summary(low_original)["confidence_meaningful"])

    untranslated = page(untranslated_es())
    check("untranslated text is headed Transcript, not English translation",
          "<h2 style=\"margin:0\">Transcript</h2>" in untranslated
          and "<h2 style=\"margin:0\">English translation</h2>" not in untranslated)
    check("it says plainly that this is not an English translation",
          "This is not an English translation." in untranslated and 'id="output-notice"' in untranslated)
    check("it names the model that failed to translate", "medium" in untranslated)
    check("it points to a way forward", "choose a standard multilingual" in " ".join(untranslated.split()))
    check("proofreading is off with a reason", 'id="proofread-language-note"' in untranslated
          and "Spanish" in untranslated and '"language_supported": false' in untranslated)
    check("its header says translation not confirmed",
          "Translation requested - the text is NOT English" in untranslated)

    mixed = build_doc(["我用", "Zoom", "开会", "and", "then", "we", "talked", "about", "the",
                       "plan", "for", "the", "week", "and", "the", "budget"],
                      language="zh", operation="translate")
    unconfirmed = page(mixed)
    check("unconfirmed output gets its own notice", "English translation not confirmed." in unconfirmed)
    check("and proofreading is off because the language is unknown",
          "could not be established" in " ".join(unconfirmed.split()) and "language: \"und\"" in unconfirmed)

    chinese = page(untranslated_zh())
    check("untranslated Chinese is handed to the editor as a no-space script",
          "spaceless: true" in chinese and 'language: "zh"' in chinese)
    english_page = page(translated_en())
    check("a real translation is handed over as spaced English",
          "spaceless: false" in english_page and 'language: "en"' in english_page)

    blank = projects.add_document(slug, "empty.wav", data=b"RIFF synthetic")["id"]
    waiting = client.get(f"/project/{slug}/document/{blank}").text
    check("an unprocessed recording uses the project's current wording",
          "Not translated yet" in waiting and "Back to the project to translate it" in waiting)
    projects.save_project(slug, {"operation": "transcribe"})
    waiting = client.get(f"/project/{slug}/document/{blank}").text
    check("and follows the project when it changes",
          "Not transcribed yet" in waiting and "Back to the project to transcribe it" in waiting)


def test_existing_records_keep_their_labels():
    main, client = _client()
    slug = new_project("Antiguos")

    # A record made before operations were stored: nothing but the engine task.
    old_original = build_doc(ES_TOKENS, language="es", operation="")
    old_original.pop("operation"); old_original.pop("output_check")
    old_translate = build_doc(ES_TOKENS, language="es", operation="", legacy_task="translate")
    old_translate.pop("operation"); old_translate.pop("output_check")
    old_english = build_doc(EN_TOKENS, language="zh", operation="", legacy_task="translate")
    old_english.pop("operation"); old_english.pop("output_check")

    ids = {}
    for name, doc in (("original", old_original), ("untranslated", old_translate), ("english", old_english)):
        doc_id = projects.add_document(slug, f"{name}.wav", data=b"RIFF synthetic")["id"]
        projects.write_json(projects.transcript_path(slug, doc_id), doc)      # as an old build wrote it
        projects.set_status(slug, doc_id, "transcribed", "done")
        check(f"{name}: the old record has no result fields", "output_status" not in projects.load_document(slug, doc_id))
        ids[name] = doc_id

    docs = {d["id"]: d for d in projects.list_documents(slug)}
    check("an old transcription is labelled a transcription",
          docs[ids["original"]]["operation"] == "transcribe" and docs[ids["original"]]["output_status"] == "original")
    check("an old translate-task record with Spanish text is NOT called a translation",
          docs[ids["untranslated"]]["operation"] == "translate"
          and docs[ids["untranslated"]]["output_status"] == "untranslated")
    check("an old translate-task record with English text is a translation",
          docs[ids["english"]]["output_status"] == "translated")
    check("the labels were saved once, additively", projects.load_document(slug, ids["untranslated"])["output_status"] == "untranslated")
    check("existing fields were left alone",
          projects.load_document(slug, ids["original"])["filename"] == "original.wav"
          and projects.load_document(slug, ids["original"])["status"] == "transcribed")

    page = client.get(f"/project/{slug}/document/{ids['untranslated']}").text
    check("an old untranslated record shows the notice", "This is not an English translation." in page)
    page = client.get(f"/project/{slug}/document/{ids['original']}").text
    check("an old original record opens normally", "<h2 style=\"margin:0\">Transcript</h2>" in page)

    out = _TMP / "legacy-out"
    exporters.export_all(old_original, out, SETTINGS.all(), {"title": "t"})
    txt = (out / "transcript.txt").read_text(encoding="utf-8")
    check("an old original exports as it always did, with an operation line added",
          "Operation: Transcription" in txt and "Language: Spanish (es)" in txt)
    check("an old record's transcript file is not rewritten by being listed",
          "operation" not in projects.read_json(projects.transcript_path(slug, ids["original"])))


# ---------------------------------------------------------------------------
# Exports
# ---------------------------------------------------------------------------

def test_export_names_headers_and_metadata():
    from docx import Document

    for label, doc, stem, other in (
        ("English translation", translated_en(), "translation-en", "transcript"),
        ("unconfirmed output", untranslated_es(), "transcript", "translation-en"),
        ("transcription", original_es(), "transcript", "translation-en"),
    ):
        out = _TMP / f"export-{stem}-{label.split()[0]}"
        out.mkdir(parents=True, exist_ok=True)
        for fmt in exporters.ALL_FORMATS:                     # a stale file from the other kind of result
            (out / f"{other}.{fmt}").write_text("stale", encoding="utf-8")
        exporters.export_all(doc, out, SETTINGS.all(), {"title": "Entrevista", "project": "Estudio"})
        names = sorted(p.name for p in out.iterdir())
        check(f"{label}: files are named {stem}.*",
              names == sorted(f"{stem}.{f}" for f in exporters.ALL_FORMATS), str(names))

    out = _TMP / "export-translation-en-English"
    txt = (out / "translation-en.txt").read_text(encoding="utf-8")
    check("translation header says what was asked", "Operation: Translate to English (Whisper translate task" in txt)
    check("and that it is a machine translation of the spoken language",
          "Chinese (zh), detected automatically (97% confidence); English machine translation" in txt)
    md = (out / "translation-en.md").read_text(encoding="utf-8")
    check("markdown front matter records the operation and result",
          'operation: "translate"' in md and 'output_status: "translated"' in md
          and 'output_language: "en"' in md and 'translated_to: "en"' in md and 'language: "zh"' in md)
    check("markdown is headed English translation and says Translated offline",
          "## English translation" in md and 'processing: "Translated offline' in md)
    data = json.loads((out / "translation-en.json").read_text(encoding="utf-8"))
    check("json carries the operation, status and both languages",
          data["summary"]["operation"] == "translate" and data["summary"]["output_status"] == "translated"
          and data["summary"]["language"] == "zh" and data["summary"]["output_language"] == "en"
          and data["transcript"]["operation"] == "translate")
    docx = Document(str(out / "translation-en.docx"))
    cells = [c.text for t in docx.tables for r in t.rows for c in r.cells]
    check("docx table lists the operation and the language story",
          any(c.startswith("Translate to English") for c in cells)
          and any("English machine translation" in c for c in cells))
    check("docx body is headed English translation",
          any(p.text == "English translation" for p in docx.paragraphs))

    out = _TMP / "export-transcript-unconfirmed"
    txt = (out / "transcript.txt").read_text(encoding="utf-8")
    check("a failed translation says so in the header",
          "This is NOT an English translation" in txt and "Operation: Translate to English" in txt)
    md = (out / "transcript.md").read_text(encoding="utf-8")
    check("and in the metadata, without claiming translated_to",
          'output_status: "untranslated"' in md and "translated_to" not in md and "## Transcript" in md)

    out = _TMP / "export-transcript-transcription"
    txt = (out / "transcript.txt").read_text(encoding="utf-8")
    check("a transcription says so", "Operation: Transcription (text in the language spoken)" in txt
          and "translat" not in txt.lower())
    md = (out / "transcript.md").read_text(encoding="utf-8")
    check("with the original metadata keys intact",
          'language: "es"' in md and 'language_source: "detected"' in md and 'operation: "transcribe"' in md
          and 'processing: "Transcribed offline' in md)


# ---------------------------------------------------------------------------
# Replacing a result needs consent, and keeps both
# ---------------------------------------------------------------------------

def test_replacing_a_result_needs_consent():
    main, client = _client()
    SETTINGS.update({"model": "medium", "language": "auto", "task": "transcribe"})
    slug = new_project("Consentimiento")
    doc = original_es()
    doc["human_edited"] = True
    doc["segments"][0]["text"] = "Texto corregido por una persona."
    doc_id = store(slug, doc)
    original_bytes = projects.transcript_path(slug, doc_id).read_bytes()
    url = f"/api/projects/{slug}/transcribe"
    pending = main.QUEUE.pending_count()

    with Patched((main.ENGINE, "status", status_for("medium")),
                 (main.ENGINE, "hardware", lambda refresh=False: HW_CPU8)):
        r = client.post(url, json={"documents": [doc_id], "force": True, "operation": "translate"})
        body = r.json()
        check("Translate on a finished transcript is not queued without consent",
              r.status_code == 200 and body["queued"] == [] and main.QUEUE.pending_count() == pending, r.text)
        check("it explains what is needed", body["skipped"][0]["needs_confirmation"] is True
              and "already has a transcript" in body["skipped"][0]["reason"]
              and "English translation" in body["skipped"][0]["reason"]
              and "Edit history" in body["skipped"][0]["reason"], str(body["skipped"]))
        check("and the transcript is untouched",
              projects.transcript_path(slug, doc_id).read_bytes() == original_bytes)

        r = client.post(url, json={"operation": "translate"})
        check("the project-wide button never replaces finished results",
              r.json()["queued"] == [] and "Already transcribed" in r.json()["skipped"][0]["reason"], r.text[:200])
        r = client.post(url, json={"documents": [doc_id], "operation": "transcribe"})
        check("Redo of the same operation also needs consent", r.json()["queued"] == [])

        r = client.post(url, json={"documents": [doc_id], "force": True, "replace": True, "operation": "translate"})
        check("with consent it is queued", len(r.json()["queued"]) == 1, r.text)
        check("and marked as being translated",
              projects.load_document(slug, doc_id)["pending_operation"] == "translate")
        check("while the existing transcript is still there until the new one is ready",
              projects.transcript_path(slug, doc_id).read_bytes() == original_bytes)

    old_outputs = outputs(slug, doc_id)
    check("the old exports are there before the run", all(n.startswith("transcript.") for n in old_outputs))

    stub, error = run_job(slug, doc_id, operation="translate", texts=EN_JOB_TEXTS, detected="es")
    check("the run succeeded", error is None, str(error))
    now = projects.load_transcript(slug, doc_id)
    check("the current result is now the translation", now["operation"] == "translate"
          and projects.load_document(slug, doc_id)["output_status"] == "translated")
    check("the pending marker is cleared", projects.load_document(slug, doc_id)["pending_operation"] is None)

    kept = [r for r in projects.list_revisions(slug, doc_id) if r["kept"]]
    check("the earlier transcript was kept", len(kept) == 1 and kept[0]["operation"] == "transcribe", str(kept))
    preserved = projects.load_revision(slug, doc_id, kept[0]["id"])
    check("with its human edits", preserved["human_edited"] is True
          and preserved["segments"][0]["text"] == "Texto corregido por una persona.")
    check("exports now describe the translation only",
          all(n.startswith("translation-en.") for n in outputs(slug, doc_id)), str(outputs(slug, doc_id)))

    # The kept copy survives ordinary autosave pruning.
    for i in range(6):
        projects.save_transcript(slug, doc_id, now, revision=True, keep=2)
    check("autosave pruning never removes a kept result",
          any(r["id"] == kept[0]["id"] for r in projects.list_revisions(slug, doc_id)))
    check("while ordinary revisions are still pruned",
          len([r for r in projects.list_revisions(slug, doc_id) if not r["kept"]]) <= 2)

    # Restoring is a second consented step, and keeps the translation too.
    r = client.post(f"/api/projects/{slug}/documents/{doc_id}/revisions/{kept[0]['id']}/restore")
    check("restoring the original works", r.status_code == 200)
    back = projects.load_transcript(slug, doc_id)
    check("the original transcript is current again", back["operation"] == "transcribe"
          and back["segments"][0]["text"] == "Texto corregido por una persona.")
    check("its record and exports follow it",
          projects.load_document(slug, doc_id)["output_status"] == "original"
          and all(n.startswith("transcript.") for n in outputs(slug, doc_id)), str(outputs(slug, doc_id)))
    kept_ops = sorted(r["operation"] for r in projects.list_revisions(slug, doc_id) if r["kept"])
    check("and the translation is kept for later", "translate" in kept_ops, str(kept_ops))

    page = client.get(f"/project/{slug}/document/{doc_id}").text
    check("Edit history labels the earlier results", "earlier English translation" in page)


def test_a_failed_or_cancelled_run_keeps_the_existing_result():
    SETTINGS.update({"model": "medium"})
    slug = new_project("Fallos")
    doc_id = store(slug, original_es())
    before = projects.transcript_path(slug, doc_id).read_bytes()
    files = outputs(slug, doc_id)

    stub, error = run_job(slug, doc_id, operation="translate", texts=[], raises=RuntimeError("model blew up"))
    check("a failed run reports the failure", error == "model blew up")
    check("the transcript is untouched", projects.transcript_path(slug, doc_id).read_bytes() == before)
    check("its exports are untouched", outputs(slug, doc_id) == files)
    check("nothing was archived for a run that produced nothing",
          not [r for r in projects.list_revisions(slug, doc_id) if r["kept"]])
    check("the failure is shown on the recording", projects.load_document(slug, doc_id)["status"] == "failed")

    fresh = projects.add_document(slug, "n.wav", data=b"RIFF synthetic")["id"]
    run_job(slug, fresh, operation="transcribe", texts=ES_JOB_TEXTS)
    check("a first result needs no archive", not [r for r in projects.list_revisions(slug, fresh) if r["kept"]])


# ---------------------------------------------------------------------------
# Proofreading through the API
# ---------------------------------------------------------------------------

def test_proofreading_api_uses_the_actual_text():
    main, client = _client()
    SETTINGS.update({"proofread_spelling": True, "proofread_artefacts": True, "task": "transcribe"})
    slug = new_project("Revision")
    body = {"segments": [{"id": 0, "text": "Habla habla con con calma, señora Núñez."}]}

    es = store(slug, untranslated_es())
    r = client.post(f"/api/projects/{slug}/documents/{es}/proofread", json=body).json()
    check("Spanish returned for a translation gets no English flags",
          r["enabled"] is False and r["issues"] == {} and r["reason"] == "language"
          and r["language_name"] == "Spanish", str(r))

    mixed = store(slug, build_doc(["我用", "Zoom", "开会", "and", "then", "we", "talked", "about", "the",
                                   "plan", "for", "the", "week", "and", "the", "budget"],
                                  language="zh", operation="translate"))
    r = client.post(f"/api/projects/{slug}/documents/{mixed}/proofread", json=body).json()
    check("unconfirmed output gets no flags either", r["enabled"] is False and r["language"] == "und", str(r))

    en = store(slug, translated_en())
    r = client.post(f"/api/projects/{slug}/documents/{en}/proofread",
                    json={"segments": [{"id": 0, "text": "I saw the the dog."}]}).json()
    check("a real English translation is proofread", r["enabled"] is True and len(r["issues"]["0"]) > 0, str(r))

    original = store(slug, build_doc(EN_TOKENS, language="en", operation="transcribe"))
    r = client.post(f"/api/projects/{slug}/documents/{original}/proofread",
                    json={"segments": [{"id": 0, "text": "I saw the the dog."}]}).json()
    check("English transcripts are proofread as before", r["enabled"] is True and len(r["issues"]["0"]) > 0)


# ---------------------------------------------------------------------------
# Settings page and the browser-facing code
# ---------------------------------------------------------------------------

def test_settings_page_and_scripts():
    main, client = _client()
    SETTINGS.update({"task": "transcribe", "model": "large-v3-turbo"})
    with Patched((main.ENGINE, "hardware", lambda refresh=False: HW_CPU8)):
        page = client.get("/settings").text
    flat = " ".join(page.split())
    check("the settings page recommends a model for transcription and another for translation",
          "Recommended: small" in flat and "For Translate to English, recommended: small" in flat, flat[:0])
    check("the translation recommendation has its own button", 'id="apply-rec-translate"' in page
          and 'id="apply-rec"' in page)
    check("the model table has a Translate to English column", "<th>Translate to English</th>" in page)
    check("Turbo is marked unsuitable in the table",
          re.search(r'data-model="large-v3-turbo".*?data-translation="unsupported"', page, re.S) is not None)
    check("Tiny is marked as poor quality", re.search(r'data-model="tiny".*?data-translation="limited"', page, re.S) is not None)
    check("the default operation is labelled for new projects only",
          "Default operation for new projects" in page and "Translate to English</option>" in page
          and "never alters existing projects" in flat)

    js = (ROOT / "app" / "static" / "js" / "project.js").read_text(encoding="utf-8")
    check("the project script takes its wording from the server's terms",
          "info.terms" in js and "T(operation).action" in js)
    check("and sends the operation with every run", "language, operation" in js)
    check("it asks before replacing a finished result", "will replace it as the current result" in js
          and "Edit history" in js)
    posts = re.findall(r"LS\.post\('/api/settings'.*?\)", js, re.S)
    check("the only settings write is the explicit Use-this-model button, and only for the model",
          len(posts) == 1 and "model: suggestion.key" in posts[0]
          and "operation" not in posts[0] and "task" not in posts[0], str(posts))
    check("per-recording buttons are labelled for screen readers", 'aria-label="${LS.escapeHtml(buttonLabel)}' in js)


def main() -> int:
    tests = [
        test_english_output_check,
        test_output_status_and_language,
        test_text_rules_follow_the_actual_text,
        test_model_classification,
        test_recommendations,
        test_operation_default_and_persistence,
        test_switching_operation_does_not_touch_defaults_or_documents,
        test_api_enforces_model_and_operation,
        test_api_operation_resolution_and_errors,
        test_worker_refuses_a_model_that_changed_while_queued,
        test_global_default_does_not_leak_into_a_run,
        test_job_records_what_happened,
        test_terms_table,
        test_project_page_follows_the_operation,
        test_document_page_labels_follow_the_document,
        test_existing_records_keep_their_labels,
        test_export_names_headers_and_metadata,
        test_replacing_a_result_needs_consent,
        test_a_failed_or_cancelled_run_keeps_the_existing_result,
        test_proofreading_api_uses_the_actual_text,
        test_settings_page_and_scripts,
    ]
    try:
        for test in tests:
            print(f"\n{test.__name__}")
            test()
    finally:
        shutil.rmtree(_TMP, ignore_errors=True)

    print()
    if FAILURES:
        print(f"{len(FAILURES)} check(s) failed: {FAILURES}")
        return 1
    print("all checks passed")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
