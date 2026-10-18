"""Tests for multilingual transcription (original-language, Phase 1).

Covers, with synthetic data only:

* language codes, names, validation and per-model support
* the saved language setting and its migration from free text
* the parameters handed to the Whisper engines (auto vs explicit)
* the language surviving the job pipeline and a save/load round trip
* Unicode in every export, and captions for scripts written without spaces
* re-timestamping such scripts
* proofreading being switched off for non-English transcripts
* the HTTP API and pages (needs httpx: pip install -r requirements-dev.txt)

Nothing here loads a Whisper model or touches real audio: engine behaviour is
checked against stub models, so these tests show that the right language is
*requested* and *recorded*, not that a given recording transcribes well.

All state lives in a throwaway data directory; no real project is read.

Run with:  .venv/Scripts/python.exe -m tests.test_languages
"""

from __future__ import annotations

import json
import os
import shutil
import sys
import tempfile
from pathlib import Path
from types import SimpleNamespace

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

# Everything the app persists must go to a temp folder, set before app.config
# is imported because the paths are fixed at import time.
_TMP = Path(tempfile.mkdtemp(prefix="localscribe-lang-test-")).resolve()
os.environ["LOCALSCRIBE_DATA_DIR"] = str(_TMP)

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except (AttributeError, ValueError):
    pass

from app import engine as E                                   # noqa: E402
from app import engine_mlx, exporters, jobs, languages, projects, proofread, realign  # noqa: E402
from app.config import PATHS, SETTINGS, SettingsStore, SPECS, coerce_value  # noqa: E402
from app.transcript import (                                  # noqa: E402
    make_segment, make_word, new_document, normalise, summary, word_count,
)

assert str(PATHS.projects).startswith(str(_TMP)), "tests must not touch real data"

FAILURES: list = []


def check(label: str, condition: bool, detail: str = "") -> None:
    if condition:
        print(f"  ok   {label}")
    else:
        print(f"  FAIL {label} {ascii(detail)}")
        FAILURES.append(label)


# ---------------------------------------------------------------------------
# Synthetic documents
# ---------------------------------------------------------------------------

def timed(tokens, start=0.0, step=0.4):
    words, cursor = [], start
    for token in tokens:
        words.append(make_word(token, cursor, cursor + step - 0.05, 0.9))
        cursor += step
    return words


def spanish_doc(requested="auto"):
    doc = new_document(duration=8.0, model="large-v3-turbo", language="es",
                       language_probability=0.9712, language_requested=requested)
    first = timed(["¿Cómo", "está", "usted,", "señora", "Núñez?"])
    second = timed(["Muy", "bien,", "gracias.", "Mañana", "vamos", "al", "niño."], start=3.0)
    doc["segments"] = [
        make_segment(0, 0.0, first[-1]["end"], "¿Cómo está usted, señora Núñez?",
                     first, speaker="José"),
        make_segment(1, 3.0, second[-1]["end"], "Muy bien, gracias. Mañana vamos al niño.",
                     second, speaker="Ana María"),
    ]
    doc["speakers"] = []
    normalise(doc)
    return doc


def chinese_doc():
    doc = new_document(duration=8.0, model="large-v3", language="zh",
                       language_probability=0.99, language_requested="zh")
    tokens = ["今天", "天气", "很", "好。", "我", "用", "Zoom", "开会。"]
    words = timed(tokens)
    doc["segments"] = [
        make_segment(0, 0.0, words[-1]["end"], "今天天气很好。我用 Zoom 开会。", words),
    ]
    normalise(doc)
    return doc


def stub_settings(**values) -> SettingsStore:
    store = SettingsStore(_TMP / f"settings-{len(list(_TMP.glob('settings-*')))}.json")
    if values:
        store.update(values)
    return store


# ---------------------------------------------------------------------------
# Language table and validation
# ---------------------------------------------------------------------------

def test_parse_and_names():
    check("empty means auto", languages.parse(None) == "auto" and languages.parse("") == "auto")
    check("auto spellings", {languages.parse(v) for v in ("auto", "AUTO", " Auto ", "detect")} == {"auto"})
    check("codes", languages.parse("es") == "es" and languages.parse("ES") == "es")
    check("names", languages.parse("Spanish") == "es" and languages.parse("mandarin") == "zh")
    check("regional tags", languages.parse("pt-BR") == "pt" and languages.parse("zh_CN") == "zh")
    check("unknown is None", languages.parse("klingon") is None and languages.parse("e") is None)
    check("coerce falls back to auto", languages.coerce("klingon") == "auto")
    check("name", languages.name("es") == "Spanish" and languages.name("") == "")
    check("label", languages.label("es") == "Spanish (es)" and languages.label("") == "")
    check("auto name", languages.name("auto") == "Automatic detection")


def test_choices():
    opts = languages.choices()
    values = [o["value"] for o in opts]
    check("automatic is first", values[0] == "auto")
    check("every language offered once", len(values) == len(set(values)) == 1 + len(languages.LANGUAGES))
    check("common languages present", {"en", "es", "zh", "fr", "ar", "hi"} <= set(values))
    labels = [o["label"] for o in opts[1:]]
    check("sorted by name", labels == sorted(labels))


def test_codes_match_whisper_tokenizer():
    # faster-whisper is a runtime requirement, so a failed import is a broken
    # environment, not a reason to skip: the offered codes must be ones the
    # engine's tokenizer actually accepts.
    from faster_whisper.tokenizer import _LANGUAGE_CODES
    missing = set(languages.LANGUAGES) - set(_LANGUAGE_CODES)
    check("every offered code is a Whisper code", not missing, str(sorted(missing)))
    check("spaceless set is a subset", languages.SPACELESS <= set(languages.LANGUAGES))
    check("right-to-left set is a subset", languages.RTL <= set(languages.LANGUAGES))


def test_model_support():
    check("auto always allowed",
          languages.unsupported_reason("auto", english_only=True) is None)
    check("english on an English-only model",
          languages.unsupported_reason("en", english_only=True, model_key="tiny.en") is None)
    reason = languages.unsupported_reason("es", english_only=True, model_key="tiny.en")
    check("non-English refused on an English-only model", bool(reason) and "only transcribes English" in reason)
    check("multilingual model accepts Spanish",
          languages.unsupported_reason("es", model_key="large-v2") is None)
    check("cantonese refused on a pre-v3 model",
          bool(languages.unsupported_reason("yue", model_key="large-v2")))
    check("cantonese accepted on v3 and turbo",
          languages.unsupported_reason("yue", model_key="large-v3") is None
          and languages.unsupported_reason("yue", model_key="large-v3-turbo-mlx") is None)


def test_join_tokens():
    check("spaced languages join with spaces", languages.join_tokens(["Hola", "mundo"]) == "Hola mundo")
    check("spaceless join without spaces", languages.join_tokens(["今天", "天气"], True) == "今天天气")
    check("latin runs keep their space when spacing is unknown",
          languages.join_tokens(["see", "Zoom", "2", "好"], True) == "see Zoom 2好")


def test_word_gaps_preserve_original_spacing():
    tokens = ["我", "用", "Zoom", "开会。"]
    gaps = languages.word_gaps("我用 Zoom 开会。", tokens)
    check("spaces around an embedded word are found", gaps == [False, False, True, True], str(gaps))
    check("and used when joining",
          languages.join_tokens(tokens, True, gaps) == "我用 Zoom 开会。")

    # No space in the original means none is invented, even between Latin tokens.
    tokens = ["用", "iPhone", "15", "手机"]
    gaps = languages.word_gaps("用iPhone15手机", tokens)
    check("no space is invented", languages.join_tokens(tokens, True, gaps) == "用iPhone15手机")

    gaps = languages.word_gaps("你好， 世界", ["你好，", "世界"])
    check("a space after punctuation is kept", gaps == [False, True], str(gaps))

    gaps = languages.word_gaps("完全不同的文字", ["我", "Zoom", "Teams"])
    check("text that no longer matches falls back to the Latin rule", gaps == [False, False, True], str(gaps))


# ---------------------------------------------------------------------------
# Settings
# ---------------------------------------------------------------------------

def test_setting_schema_and_migration():
    spec = SPECS["language"]
    check("language is a dropdown", spec["kind"] == "language_select")
    check("defaults to automatic", spec["default"] == "auto")
    check("choices attached", spec["choices"][0]["value"] == "auto")
    check("legacy name migrates", coerce_value("language", "Spanish") == "es")
    check("legacy upper-case code migrates", coerce_value("language", "FR") == "fr")
    check("regional tag migrates", coerce_value("language", "pt-BR") == "pt")
    check("garbage becomes auto", coerce_value("language", "xx-not-real") == "auto")
    check("None becomes auto", coerce_value("language", None) == "auto")

    path = _TMP / "legacy-settings.json"
    path.write_text(json.dumps({"language": "German", "model": "small"}), encoding="utf-8")
    store = SettingsStore(path)
    check("a saved free-text value loads as a code", store.get("language") == "de")
    store.update({"language": "ja"})
    check("a saved value round-trips", SettingsStore(path).get("language") == "ja")


# ---------------------------------------------------------------------------
# Engine parameters
# ---------------------------------------------------------------------------

def params_for(settings: SettingsStore, overrides=None) -> dict:
    original = E.SETTINGS
    E.SETTINGS = settings
    try:
        return E.WhisperEngine().transcribe_params(overrides)
    finally:
        E.SETTINGS = original


def test_engine_params():
    check("default sends no language (detect)", "language" not in params_for(stub_settings()))
    check("saved language is sent", params_for(stub_settings(language="es"))["language"] == "es")
    check("override beats the saved language",
          params_for(stub_settings(language="es"), {"language": "fr"})["language"] == "fr")
    check("auto override removes a saved language",
          "language" not in params_for(stub_settings(language="es"), {"language": "auto"}))
    check("names are normalised to codes",
          params_for(stub_settings(), {"language": "German"})["language"] == "de")
    check("regional tags are normalised",
          params_for(stub_settings(), {"language": "pt-BR"})["language"] == "pt")

    advanced = stub_settings()
    advanced.update({"advanced_raw": {"language": "it"}})
    check("advanced passthrough still works", params_for(advanced)["language"] == "it")
    check("an explicit choice beats advanced",
          params_for(advanced, {"language": "es"})["language"] == "es")
    check("auto beats advanced", "language" not in params_for(advanced, {"language": "auto"}))

    p = params_for(stub_settings(), {"language": "ja"})
    check("existing parameters are untouched",
          p["task"] == "transcribe" and p["word_timestamps"] is True and p["beam_size"] == 5)


def test_mlx_params():
    forced = engine_mlx._translate_params({"language": "es", "task": "transcribe", "beam_size": 5})
    check("MLX receives a forced language", forced.get("language") == "es")
    detect = engine_mlx._translate_params({"task": "transcribe", "beam_size": 5})
    check("MLX omits language for detection", "language" not in detect)


class FakeModel:
    def __init__(self, info):
        self.info = info
        self.calls = []

    def transcribe(self, audio, **params):
        self.calls.append(params)
        word = SimpleNamespace(word=" Hola", start=0.0, end=0.5, probability=0.9)
        seg = SimpleNamespace(start=0.0, end=1.0, text=" Hola", words=[word], tokens=[1],
                              avg_logprob=-0.2, no_speech_prob=0.01,
                              compression_ratio=1.1, temperature=0.0)
        return iter([seg]), self.info


def run_engine(settings, info, overrides=None):
    model = FakeModel(info)
    eng = E.WhisperEngine()
    eng._backend = "ctranslate2"
    eng.ensure_loaded = lambda: model
    original = E.SETTINGS
    E.SETTINGS = settings
    try:
        result = eng._transcribe_once([0.0] * 16000, duration=1.0, overrides=overrides)
    finally:
        E.SETTINGS = original
    return model, result


def test_engine_reports_language():
    info = SimpleNamespace(language="es", language_probability=0.93, duration_after_vad=None)
    model, result = run_engine(stub_settings(), info)
    check("detection: no language passed to the model", "language" not in model.calls[0])
    check("detection: detected language reported", result["language"] == "es")
    check("detection: confidence reported", abs(result["language_probability"] - 0.93) < 1e-6)
    check("detection: recorded as auto", result["language_requested"] == "auto")

    info = SimpleNamespace(language="fr", language_probability=1.0, duration_after_vad=None)
    model, result = run_engine(stub_settings(), info, {"language": "fr"})
    check("explicit: language passed to the model", model.calls[0].get("language") == "fr")
    check("explicit: recorded as requested", result["language_requested"] == "fr")
    check("explicit: language in the logged parameters", result["params"].get("language") == "fr")

    silent = SimpleNamespace(language="", language_probability=0.0, duration_after_vad=None)
    _, result = run_engine(stub_settings(), silent, {"language": "de"})
    check("explicit: language kept if the back end reports none", result["language"] == "de")


# ---------------------------------------------------------------------------
# Persistence
# ---------------------------------------------------------------------------

def test_language_persists():
    doc = spanish_doc("auto")
    path = _TMP / "roundtrip.json"
    projects.write_json(path, doc)
    loaded = projects.read_json(path)
    check("language survives a save/load", loaded["language"] == "es")
    check("confidence survives", abs(loaded["language_probability"] - 0.9712) < 1e-9)
    check("request survives", loaded["language_requested"] == "auto")
    check("non-ASCII text is stored readably", "¿Cómo" in path.read_text(encoding="utf-8"))

    info = summary(loaded)
    check("summary has language fields",
          info["language"] == "es" and info["language_name"] == "Spanish"
          and info["language_requested"] == "auto" and info["output_language"] == "es")


def test_legacy_documents_still_load():
    legacy = new_document(duration=3.0, model="small", language="en", language_probability=0.99)
    legacy.pop("language_requested")
    legacy["segments"] = [make_segment(0, 0.0, 2.0, "Hello there.", timed(["Hello", "there."]))]
    normalise(legacy)
    info = summary(legacy)
    check("legacy summary works", info["language"] == "en" and info["language_requested"] == "")
    check("legacy language stays English", info["output_language"] == "en")
    check("legacy exports have a language line",
          "Language: English (en)" in "\n".join(exporters._plain_header(legacy, {})))

    no_language = new_document(duration=1.0, model="small")
    check("a document with no language is tolerated", summary(no_language)["language"] == "")
    check("and writes no language line",
          not any(r.startswith("Language") for r in exporters._plain_header(no_language, {})))


def test_word_count_by_script():
    check("English words counted as before",
          word_count(spanish_doc()) == 5 + 7)
    zh = chinese_doc()
    check("Chinese counts characters, not one token", word_count(zh) == 14, str(word_count(zh)))
    check("Chinese summary flags the unit", summary(zh)["spaceless"] is True)


# ---------------------------------------------------------------------------
# Job pipeline
# ---------------------------------------------------------------------------

class StubEngine:
    """Stands in for the Whisper engine inside the job runner."""

    def __init__(self, detected="es", probability=0.9712, task="transcribe"):
        self.detected = detected
        self.probability = probability
        self.task = task
        self.overrides = None

    def transcribe(self, pcm, *, duration, overrides=None, on_segment=None, should_cancel=None):
        self.overrides = dict(overrides or {})
        forced = self.overrides.get("language")
        language = forced if forced and forced != "auto" else self.detected
        tracker = SimpleNamespace(snapshot=lambda: {
            "progress": 0.5, "eta_text": "1s", "words_per_second": 2.0,
            "tokens_per_second": 4.0, "realtime_factor": 3.0,
        })
        texts = [" ¿Cómo está usted, señora Núñez?", " Muy bien, gracias."]
        segments = []
        for index, text in enumerate(texts):
            tokens = text.split()
            words = [SimpleNamespace(word=f" {t}", start=index * 3 + i * 0.4,
                                     end=index * 3 + i * 0.4 + 0.3, probability=0.9)
                     for i, t in enumerate(tokens)]
            seg = SimpleNamespace(start=index * 3.0, end=index * 3.0 + 2.5, text=text,
                                  words=words, tokens=[1, 2], avg_logprob=-0.2,
                                  no_speech_prob=0.01, compression_ratio=1.1, temperature=0.0)
            segments.append(seg)
            on_segment(index, seg, tracker)
        return {
            "segments": segments,
            "info": SimpleNamespace(language=language),
            "language": language,
            "language_probability": 1.0 if forced and forced != "auto" else self.probability,
            "language_requested": forced if forced and forced != "auto" else "auto",
            "stats": {"wall_seconds": 1.0, "realtime_factor": 3.0, "model": "stub",
                      "device": "cpu", "compute_type": "int8"},
            "params": {"task": self.task, **({"language": forced} if forced not in (None, "auto") else {})},
        }


def run_job(slug, doc_id, language, detected="es", task="transcribe"):
    stub = StubEngine(detected=detected, task=task)
    saved = (jobs.ENGINE, jobs.media.probe, jobs.media.decode_to_pcm,
             jobs.media.write_waveform_cache, jobs.media.describe)
    info = SimpleNamespace(
        container="wav", audio_codec="pcm", sample_rate=16000, channels=1, duration=6.0,
        kind="audio", browser_playable=True, has_video=False, video_codec=None,
        width=None, height=None, to_dict=lambda: {"container": "wav", "has_video": False},
    )
    jobs.ENGINE = stub
    jobs.media.probe = lambda *a, **k: info
    jobs.media.decode_to_pcm = lambda *a, **k: [0.0] * (16000 * 6)
    jobs.media.write_waveform_cache = lambda *a, **k: None
    jobs.media.describe = lambda i: "synthetic"
    try:
        job = jobs.Job("transcribe", "entrevista.wav", slug=slug, doc_id=doc_id,
                       language=language, operation=task)
        jobs.JobQueue()._do_transcription(job)
    finally:
        (jobs.ENGINE, jobs.media.probe, jobs.media.decode_to_pcm,
         jobs.media.write_waveform_cache, jobs.media.describe) = saved
    return stub


def test_job_queue_carries_language():
    queue = jobs.JobQueue()
    job = queue.submit_transcription("p", "d", "file.wav", language="es")
    check("language is carried on the job", job.payload["language"] == "es")
    job = queue.submit_transcription("p", "d", "file.wav")
    check("no language leaves settings in charge", job.payload["language"] is None)


def test_language_flows_through_the_job():
    project = projects.create_project("Entrevistas")
    slug = project["slug"]
    record = projects.add_document(slug, "entrevista.wav", data=b"RIFF synthetic")
    doc_id = record["id"]

    stub = run_job(slug, doc_id, "es")
    check("explicit language reaches the engine", stub.overrides.get("language") == "es")
    doc = projects.load_transcript(slug, doc_id)
    check("saved transcript records language", doc["language"] == "es")
    check("saved transcript records the request", doc["language_requested"] == "es")
    check("segments saved with accents", doc["segments"][0]["text"].startswith("¿Cómo"))
    stored = projects.load_document(slug, doc_id)
    check("document record carries the language",
          stored["language"] == "es" and stored["language_name"] == "Spanish"
          and stored["language_requested"] == "es")

    out = projects.outputs_dir(slug, doc_id)
    txt = (out / "transcript.txt").read_text(encoding="utf-8")
    check("export names a selected language", "Language: Spanish (es), selected by the user" in txt)
    check("export text is Unicode intact", "¿Cómo está usted, señora Núñez?" in txt)

    stub = run_job(slug, doc_id, "auto", detected="es")
    check("auto reaches the engine as auto", stub.overrides.get("language") == "auto")
    doc = projects.load_transcript(slug, doc_id)
    check("detected language is stored", doc["language"] == "es" and doc["language_requested"] == "auto")
    check("detection confidence is stored", abs(doc["language_probability"] - 0.9712) < 1e-6)
    txt = (out / "transcript.txt").read_text(encoding="utf-8")
    check("export says it was detected, with confidence",
          "Language: Spanish (es), detected automatically (97% confidence)" in txt)

    stub = run_job(slug, doc_id, None)
    check("no request leaves the engine to its settings", "language" not in stub.overrides)


# ---------------------------------------------------------------------------
# Unicode exports and captions
# ---------------------------------------------------------------------------

def export_to(doc, name, **settings):
    from app.config import default_settings
    merged = default_settings()
    merged.update(settings)
    out = _TMP / name
    written = exporters.export_all(doc, out, merged, {"title": "Entrevista", "project": "Estudio"})
    return out, written


def test_unicode_exports():
    out, written = export_to(spanish_doc("auto"), "es-out")
    check("every format wrote", not any(str(v).startswith("error") for v in written.values()), str(written))

    txt = (out / "transcript.txt").read_text(encoding="utf-8")
    check("txt keeps accents and inverted marks", "¿Cómo está usted, señora Núñez?" in txt)
    check("txt keeps accented speaker names", "José:" in txt and "Ana María:" in txt)

    md = (out / "transcript.md").read_text(encoding="utf-8")
    check("md front matter has code and name", 'language: "es"' in md and 'language_name: "Spanish"' in md)
    check("md records detection", 'language_source: "detected"' in md and "language_confidence: 0.9712" in md)
    check("md body keeps accents", "Mañana vamos al niño." in md)

    for name in ("vtt", "srt"):
        text = (out / f"transcript.{name}").read_text(encoding="utf-8")
        check(f"{name} keeps accents", "¿Cómo" in text and "Núñez?" in text)
    check("vtt is still a vtt", (out / "transcript.vtt").read_text(encoding="utf-8").startswith("WEBVTT"))

    raw = (out / "transcript.json").read_text(encoding="utf-8")
    check("json is not ASCII-escaped", "¿Cómo" in raw and "\\u00bf" not in raw)
    data = json.loads(raw)
    check("json carries language fields",
          data["summary"]["language"] == "es" and data["transcript"]["language_requested"] == "auto")

    from docx import Document
    paragraphs = [p.text for p in Document(str(out / "transcript.docx")).paragraphs]
    check("docx keeps accents", any("¿Cómo está usted, señora Núñez?" in p for p in paragraphs))
    cells = [c.text for t in Document(str(out / "transcript.docx")).tables for r in t.rows for c in r.cells]
    check("docx header names the language", any(c.startswith("Spanish (es), detected automatically") for c in cells))


def test_chinese_exports_and_captions():
    zh = chinese_doc()
    cues = exporters.build_cues(zh, max_chars=42, max_lines=2, max_duration=6.0)
    check("one cue holds the whole sentence", len(cues) == 1, str(cues))
    check("cue text equals the transcript text, spaces included",
          cues[0]["text"] == "今天天气很好。我用 Zoom 开会。", cues[0]["text"])

    out, written = export_to(zh, "zh-out", vtt_max_chars=10)
    check("every format wrote", not any(str(v).startswith("error") for v in written.values()), str(written))
    vtt = (out / "transcript.vtt").read_text(encoding="utf-8")
    body = [l for l in vtt.splitlines() if l and "-->" not in l and not l.isdigit()
            and l != "WEBVTT" and not l.startswith("NOTE")]
    check("caption lines respect the width", all(len(l) <= 10 for l in body), str(body))
    check("no caption line starts with punctuation",
          not any(l[0] in "。，、！？" for l in body), str(body))
    check("captions lose no characters", "".join(body).replace(" ", "") == "今天天气很好。我用Zoom开会。", str(body))
    check("Latin word is never split across lines", any("Zoom" in l for l in body), str(body))

    txt = (out / "transcript.txt").read_text(encoding="utf-8")
    check("txt header says characters, not words", "Characters: 14" in txt and "Words:" not in txt, txt[:300])
    check("txt header names Chinese as selected", "Chinese (zh), selected by the user" in txt)
    check("txt keeps the original text", "今天天气很好。我用 Zoom 开会。" in txt)
    check("md flags the length unit", 'length_unit: "characters"' in (out / "transcript.md").read_text(encoding="utf-8"))


def test_wrapping_spaceless_text():
    lines = exporters.wrap_lines("今天天气很好我们去公园散步吧", 5, spaceless=True)
    check("breaks between characters", lines == ["今天天气很", "好我们去公", "园散步吧"], str(lines))
    lines = exporters.wrap_lines("我用Zoom开会", 6, spaceless=True)
    check("keeps Latin words whole", "Zoom" in "".join(lines) and all("Zo" not in l or "Zoom" in l for l in lines),
          str(lines))
    check("spaced wrapping is unchanged", exporters.wrap_lines("aa bb cc", 5) == ["aa bb", "cc"])
    check("fits honours spaceless", exporters.fits("今天天气很好我们去公", 5, 2, spaceless=True) is True
          and exporters.fits("今天天气很好我们去公园", 5, 2, spaceless=True) is False)


def test_retimestamp_spaceless():
    # Whisper words for Chinese can span several characters (here "今天").
    def build():
        doc = new_document(duration=4.0, model="large-v3", language="zh", language_requested="zh")
        words = timed(["今天", "天气", "很", "好。"], step=0.5)
        words[1]["prob"] = 0.31
        doc["segments"] = [make_segment(0, 0.0, words[-1]["end"], "今天天气很好。", words)]
        normalise(doc)
        return doc

    from app.transcript import apply_edits
    unchanged = build()
    before = json.dumps(unchanged["segments"][0]["words"], ensure_ascii=False)
    apply_edits(unchanged, [{"id": 0, "text": "今天天气很好。"}])
    report = realign.retimestamp(unchanged)
    after = json.dumps(unchanged["segments"][0]["words"], ensure_ascii=False)
    check("unchanged text keeps the model's words, timings and confidences", before == after)
    check("and reports them as preserved", report["preserved"] == 4 and report["redistributed"] == 0, str(report))

    doc = build()
    original = doc["segments"][0]["words"]
    first_start, first_end = original[0]["start"], original[0]["end"]
    last_start, last_end = original[3]["start"], original[3]["end"]
    apply_edits(doc, [{"id": 0, "text": "今天天气很棒。"}])
    report = realign.retimestamp(doc)
    new_words = doc["segments"][0]["words"]
    check("edited text is split per character", [w["w"] for w in new_words] == list("今天天气很") + ["棒。"],
          str([w["w"] for w in new_words]))
    check("unchanged characters are preserved, not redistributed",
          report["preserved"] == 5 and report["redistributed"] == 1, str(report))
    check("outer boundaries of a split word stay exact",
          new_words[0]["start"] == first_start and new_words[1]["end"] == first_end)
    check("inside a split word the timing is interpolated",
          first_start < new_words[0]["end"] < first_end and new_words[0]["end"] == new_words[1]["start"])
    check("model confidence is inherited by preserved characters",
          new_words[2]["prob"] == 0.31 and new_words[3]["prob"] == 0.31)
    check("the edited character spans the old word's time and has no confidence",
          new_words[5]["start"] == last_start and new_words[5]["end"] == last_end and new_words[5]["prob"] == 0.0)
    starts = [w["start"] for w in new_words]
    check("timings stay monotonic", starts == sorted(starts))
    check("spaced tokenising is unchanged", realign.tokenise("a bb  c") == ["a", "bb", "c"])


# ---------------------------------------------------------------------------
# Proofreading
# ---------------------------------------------------------------------------

def translated_doc():
    """A Chinese recording run through Whisper's translate task: English text."""
    doc = new_document(duration=8.0, model="large-v3", language="zh", language_probability=0.97,
                       language_requested="auto", engine={"task": "translate"})
    tokens = ["Today", "the", "weather", "is", "nice.", "I", "use", "Zoom", "for", "meetings."]
    words = timed(tokens)
    doc["segments"] = [make_segment(0, 0.0, words[-1]["end"], " ".join(tokens), words)]
    normalise(doc)
    return doc


def test_translate_task_is_not_a_source_language_transcript():
    doc = translated_doc()
    info = summary(doc)
    check("spoken language is still recorded", info["language"] == "zh")
    check("text language is English", info["output_language"] == "en")
    check("text is not treated as a no-space script", info["spaceless"] is False)
    check("words are counted as words", word_count(doc) == 10, str(word_count(doc)))

    cues = exporters.build_cues(doc)
    check("captions keep the spaces between English words",
          cues[0]["text"] == "Today the weather is nice. I use Zoom for meetings.", cues[0]["text"])

    out, written = export_to(doc, "translated-out")
    check("every format wrote", not any(str(v).startswith("error") for v in written.values()), str(written))
    txt = (out / "translation-en.txt").read_text(encoding="utf-8")
    check("header names source language and the translation",
          "Chinese (zh), detected automatically (97% confidence); "
          "English machine translation" in txt, txt[:400])
    check("header counts words", "Words: 10" in txt and "Characters" not in txt)
    md = (out / "translation-en.md").read_text(encoding="utf-8")
    check("md marks it as translated", 'translated_to: "en"' in md and 'language: "zh"' in md)

    check("proofreading treats it as English", proofread.supports_document(doc))
    from app.transcript import apply_edits
    apply_edits(doc, [{"id": 0, "text": "Today the weather is nice. I use Zoom for calls."}])
    realign.retimestamp(doc)
    check("re-timestamping tokenises English words, not characters",
          [w["w"] for w in doc["segments"][0]["words"]][-2:] == ["for", "calls."])

    # The same Chinese recording transcribed (not translated) is a no-space script.
    original = translated_doc()
    original["engine"] = {"task": "transcribe"}
    check("without the translate task it would be spaceless", summary(original)["spaceless"] is True)


def test_translate_task_survives_the_job_pipeline():
    slug = projects.create_project("Traducido")["slug"]
    doc_id = projects.add_document(slug, "x.wav", data=b"RIFF synthetic")["id"]
    # The stub returns Spanish text whatever it is asked, as a model that is
    # not trained for translation does. (A translation-capable model is
    # selected so the job guard lets the run start.)
    SETTINGS.update({"model": "medium"})
    run_job(slug, doc_id, "auto", detected="zh", task="translate")
    doc = projects.load_transcript(slug, doc_id)
    check("task is stored with the transcript", doc["engine"].get("task") == "translate")
    stored = projects.load_document(slug, doc_id)
    check("returned non-English text is not recorded as English",
          stored["language"] == "zh" and stored["output_language"] != "en"
          and stored["output_status"] == "untranslated", str(stored))
    txt = (projects.outputs_dir(slug, doc_id) / "transcript.txt").read_text(encoding="utf-8")
    check("export does not claim a translation", "NOT an English translation" in txt)


ARABIC = ["مرحبا", "بكم", "في", "المقابلة؟"]
HEBREW = ["שלום", "מה", "שלומך", "היום?"]


def rtl_doc(code, tokens):
    doc = new_document(duration=4.0, model="large-v3", language=code, language_probability=0.99,
                       language_requested="auto")
    words = timed(tokens)
    doc["segments"] = [make_segment(0, 0.0, words[-1]["end"], " ".join(tokens), words, speaker="INT")]
    normalise(doc)
    return doc


def test_right_to_left_languages():
    from docx import Document
    from docx.oxml.ns import qn

    def bidi(paragraph):
        ppr = paragraph._p.pPr
        return ppr is not None and ppr.find(qn("w:bidi")) is not None

    def rtl_runs(paragraph):
        return paragraph.runs and all(
            r._r.rPr is not None and r._r.rPr.find(qn("w:rtl")) is not None for r in paragraph.runs)

    for code, tokens in (("ar", ARABIC), ("he", HEBREW)):
        text = " ".join(tokens)
        doc = rtl_doc(code, tokens)
        check(f"{code}: counted in words", word_count(doc) == 4)
        check(f"{code}: flagged right-to-left", summary(doc)["rtl"] is True and summary(doc)["spaceless"] is False)

        out, written = export_to(doc, f"{code}-out")
        check(f"{code}: every format wrote", not any(str(v).startswith("error") for v in written.values()), str(written))
        check(f"{code}: txt keeps the text", text in (out / "transcript.txt").read_text(encoding="utf-8"))
        check(f"{code}: md keeps the text", text in (out / "transcript.md").read_text(encoding="utf-8"))
        vtt = (out / "transcript.vtt").read_text(encoding="utf-8")
        srt = (out / "transcript.srt").read_text(encoding="utf-8")
        check(f"{code}: captions keep the text and word order", text in vtt and text in srt)
        check(f"{code}: json is readable", text in (out / "transcript.json").read_text(encoding="utf-8"))

        paragraphs = Document(str(out / "transcript.docx")).paragraphs
        body = [p for p in paragraphs if text in p.text]
        check(f"{code}: docx transcript paragraph is right-to-left", len(body) == 1 and bidi(body[0]) and rtl_runs(body[0]))
        others = [p for p in paragraphs if text not in p.text]
        check(f"{code}: other docx paragraphs are untouched", not any(bidi(p) for p in others))

        check(f"{code}: proofreading is off", not proofread.supports_document(doc))

    en_out, _ = export_to(spanish_doc(), "ltr-control")
    check("left-to-right documents get no bidi markup",
          not any(bidi(p) for p in Document(str(en_out / "transcript.docx")).paragraphs))
    check("Arabic wraps on spaces like any spaced language",
          exporters.wrap_lines(" ".join(ARABIC), 12) == ["مرحبا بكم في", "المقابلة؟"])


def test_proofreading_is_english_only():
    english_slip = "I saw the the dog yesterday."
    base = proofread.check_segments([{"id": 1, "text": english_slip}], language="en")
    check("English is still checked", len(base["1"]) > 0)
    check("empty language is not assumed to be English",
          proofread.check_segments([{"id": 1, "text": english_slip}], language="") == {})

    spanish = [{"id": 1, "text": "Habla habla con con calma, señora Núñez."}]
    check("Spanish is never flagged", proofread.check_segments(spanish, language="es") == {})
    check("Chinese is never flagged",
          proofread.check_segments([{"id": 1, "text": "今天天气很好。"}], language="zh") == {})
    check("the same text is flagged when treated as English",
          len(proofread.check_segments(spanish, language="en")["1"]) > 0)

    check("supports English and normalizes locale codes",
          proofread.supports_language("en") and proofread.supports_language("EN_us"))
    check("rejects other languages", not proofread.supports_language("es")
          and not proofread.supports_language("zh")
          and not proofread.supports_language("")
          and not proofread.supports_language(None))

    check("a Spanish document is unsupported", not proofread.supports_document(spanish_doc()))
    check("an English document is supported",
          proofread.supports_document(new_document(duration=1, model="m", language="en")))
    check("a document with no language is unsupported", not proofread.supports_document({}))
    translated = spanish_doc()
    translated["engine"] = {"task": "translate"}
    check("a translate request that returned Spanish is not English, so not proofread",
          not proofread.supports_document(translated))
    check("an English translation is proofread", proofread.supports_document(translated_doc()))


# ---------------------------------------------------------------------------
# HTTP API and pages
# ---------------------------------------------------------------------------

def _client():
    """The FastAPI app and a test client, with the network-touching hook off.

    A missing httpx is an environment error to fix, not something to skip:
    the API checks are part of the acceptance evidence.
    """
    try:
        from fastapi.testclient import TestClient
    except ImportError as exc:
        raise SystemExit(
            "fastapi.testclient needs httpx. Install the test requirements: "
            "pip install -r requirements-dev.txt"
        ) from exc

    from app import update_check
    update_check.refresh_state_async = lambda: None           # never touch the network
    from app import main
    return main, TestClient(main.app)


def test_api_and_pages():
    main, client = _client()
    project = projects.create_project("Prueba API")
    slug = project["slug"]

    def new_doc(name="a.wav"):
        return projects.add_document(slug, name, data=b"RIFF synthetic")["id"]

    def reset(doc_id):
        projects.set_status(slug, doc_id, "uploaded", "")

    def fake_status(model="large-v3-turbo", downloaded=True):
        return lambda: {"model_downloaded": downloaded, "selected_model": model,
                        "state": "idle", "ready": False, "loaded_model": "",
                        "target_device": "cpu", "loaded_device": ""}

    real_status = main.ENGINE.status
    try:
        main.ENGINE.status = fake_status()
        doc_id = new_doc()
        url = f"/api/projects/{slug}/transcribe"

        r = client.post(url, json={"documents": [doc_id], "language": "klingon"})
        check("unknown language is a 400", r.status_code == 400 and "not a supported language" in r.json()["detail"])

        r = client.post(url, json={"documents": [doc_id], "language": "es"})
        queued = r.json().get("queued") or []
        check("explicit language is accepted", r.status_code == 200 and len(queued) == 1, r.text)
        job = main.QUEUE.get(queued[0]["job"]) if queued else None
        check("job carries the explicit language", job is not None and job.payload["language"] == "es")

        reset(doc_id)
        r = client.post(url, json={"documents": [doc_id], "language": "Mandarin"})
        job = main.QUEUE.get(r.json()["queued"][0]["job"])
        check("names are normalised before queueing", job.payload["language"] == "zh")

        reset(doc_id)
        r = client.post(url, json={"documents": [doc_id], "language": "auto"})
        job = main.QUEUE.get(r.json()["queued"][0]["job"])
        check("auto is passed as auto", job.payload["language"] == "auto")

        reset(doc_id)
        SETTINGS.update({"language": "fr"})
        r = client.post(url, json={"documents": [doc_id]})
        job = main.QUEUE.get(r.json()["queued"][0]["job"])
        check("omitted language uses the saved setting", job.payload["language"] == "fr")
        SETTINGS.update({"language": "auto"})

        reset(doc_id)
        main.ENGINE.status = fake_status("tiny.en")
        r = client.post(url, json={"documents": [doc_id], "language": "es"})
        check("English-only model refuses Spanish", r.status_code == 409 and "only transcribes English" in r.json()["detail"])
        reset(doc_id)
        r = client.post(url, json={"documents": [doc_id], "language": "en"})
        check("English-only model accepts English", r.status_code == 200 and r.json()["queued"])
        reset(doc_id)
        r = client.post(url, json={"documents": [doc_id], "language": "auto"})
        check("English-only model accepts automatic", r.status_code == 200 and r.json()["queued"])

        reset(doc_id)
        main.ENGINE.status = fake_status("large-v2")
        r = client.post(url, json={"documents": [doc_id], "language": "yue"})
        check("Cantonese refused on large-v2", r.status_code == 409 and "Large v3" in r.json()["detail"])

        main.ENGINE.status = fake_status(downloaded=False)
        r = client.post(url, json={"documents": [doc_id], "language": "es"})
        check("missing model is still reported", r.status_code == 409 and "not been downloaded" in r.json()["detail"])
    finally:
        main.ENGINE.status = real_status

    # -- pages -------------------------------------------------------------
    page = client.get(f"/project/{slug}")
    check("project page renders a language control",
          page.status_code == 200 and 'id="language-select"' in page.text)
    check("automatic is the default", '<option value="auto" selected>' in page.text)
    import re
    language_select = re.search(r'<select id="language-select".*?</select>', page.text, re.S).group(0)
    check("every language is offered",
          language_select.count("<option ") == 1 + len(languages.LANGUAGES))
    SETTINGS.update({"language": "es"})
    check("a project with no choice starts from the AI Settings default",
          '<option value="es" selected>' in client.get(f"/project/{slug}").text)
    SETTINGS.update({"language": "auto"})

    settings_page = client.get("/settings")
    check("settings page renders the language dropdown",
          settings_page.status_code == 200 and 'id="f-language"' in settings_page.text
          and '<option value="auto" selected>' in settings_page.text)

    def make_doc(doc):
        doc_id = new_doc("doc.wav")
        projects.save_transcript(slug, doc_id, doc, revision=False)
        projects.set_status(slug, doc_id, "transcribed", "")
        return doc_id

    es_id = make_doc(spanish_doc("auto"))
    es_page = client.get(f"/project/{slug}/document/{es_id}")
    check("Spanish document page renders", es_page.status_code == 200)
    check("page names the detected language", "Spanish (detected)" in es_page.text)
    check("page explains proofreading is off", 'id="proofread-language-note"' in es_page.text
          and "English-only" in es_page.text)
    check("editor is told proofreading is unsupported", '"language_supported": false' in es_page.text)
    check("page declares the language to the editor", 'language: "es"' in es_page.text)

    en_doc = new_document(duration=3.0, model="small", language="en", language_probability=0.99)
    en_doc["segments"] = [make_segment(0, 0.0, 2.0, "Hello there.", timed(["Hello", "there."]))]
    en_page = client.get(f"/project/{slug}/document/{make_doc(en_doc)}")
    check("English page keeps proofreading",
          'id="proofread-language-note"' not in en_page.text and '"language_supported": true' in en_page.text)

    zh_page = client.get(f"/project/{slug}/document/{make_doc(chinese_doc())}")
    check("Chinese page tells the editor to join without spaces",
          "spaceless: true" in zh_page.text and 'language: "zh"' in zh_page.text)
    check("Spanish page does not", "spaceless: false" in es_page.text)

    tr_page = client.get(f"/project/{slug}/document/{make_doc(translated_doc())}")
    check("a translated Chinese recording is treated as spaced English",
          tr_page.status_code == 200 and "spaceless: false" in tr_page.text
          and 'language: "en"' in tr_page.text)
    check("and keeps proofreading",
          'id="proofread-language-note"' not in tr_page.text and '"language_supported": true' in tr_page.text)
    check("and says what happened", "English machine translation" in " ".join(tr_page.text.split()))

    legacy = new_document(duration=3.0, model="small", language="en")
    legacy.pop("language_requested")
    legacy_page = client.get(f"/project/{slug}/document/{make_doc(legacy)}")
    check("legacy document page still renders", legacy_page.status_code == 200
          and 'id="proofread-language-note"' not in legacy_page.text)

    # -- proofread endpoint -------------------------------------------------
    body = {"segments": [{"id": 0, "text": "I saw the the dog."}]}
    r = client.post(f"/api/projects/{slug}/documents/{es_id}/proofread", json=body)
    data = r.json()
    check("proofread endpoint is off for Spanish",
          r.status_code == 200 and data["enabled"] is False and data["reason"] == "language"
          and data["issues"] == {} and data["language_name"] == "Spanish", r.text)

    en_id = make_doc(en_doc)
    r = client.post(f"/api/projects/{slug}/documents/{en_id}/proofread", json=body)
    check("proofread endpoint still works for English",
          r.json()["enabled"] is True and len(r.json()["issues"]["0"]) > 0, r.text)

    legacy_id = make_doc(legacy)
    r = client.post(f"/api/projects/{slug}/documents/{legacy_id}/proofread", json=body)
    check("proofread endpoint still works for legacy documents",
          r.json()["enabled"] is True and len(r.json()["issues"]["0"]) > 0, r.text)


def test_project_language_is_per_project():
    main, client = _client()
    SETTINGS.update({"language": "auto", "task": "transcribe"})
    a = projects.create_project("Proyecto A")["slug"]
    b = projects.create_project("Project B")["slug"]

    def selected(slug):
        text = client.get(f"/project/{slug}").text
        found = [v for v in languages.LANGUAGES if f'<option value="{v}" selected>' in text]
        return found[0] if found else ("auto" if '<option value="auto" selected>' in text else None)

    check("new projects have no language choice", projects.load_project(a)["transcription_language"] == "")

    r = client.put(f"/api/projects/{a}", json={"transcription_language": "Spanish"})
    check("a project remembers its language", r.json()["project"]["transcription_language"] == "es")
    check("without touching the AI Settings default", SETTINGS.get("language") == "auto")
    check("project A shows its own choice", selected(a) == "es")
    check("project B is unaffected", selected(b) == "auto")

    SETTINGS.update({"language": "fr"})
    check("projects without a choice follow the AI Settings default", selected(b) == "fr")
    check("a project with a choice does not", selected(a) == "es")

    client.put(f"/api/projects/{a}", json={"transcription_language": "klingon"})
    check("an invalid value is ignored", projects.load_project(a)["transcription_language"] == "es")
    client.put(f"/api/projects/{a}", json={"title": "Proyecto A renombrado"})
    check("saving other details keeps the language", projects.load_project(a)["transcription_language"] == "es")
    client.put(f"/api/projects/{a}", json={"transcription_language": "auto"})
    check("automatic is a real choice", projects.load_project(a)["transcription_language"] == "auto"
          and selected(a) == "auto")

    created = projects.create_project("Con idioma", transcription_language="German")
    check("a language can be set at creation", created["transcription_language"] == "de")
    created = projects.create_project("Con idioma malo", transcription_language="klingon")
    check("an invalid creation value is dropped", created["transcription_language"] == "")

    legacy = projects.project_dir(b) / projects.PROJECT_FILE
    record = projects.read_json(legacy)
    record.pop("transcription_language", None)
    projects.write_json(legacy, record)
    check("a project saved before this feature still loads", selected(b) == "fr")

    # Which language a run actually uses: request, then project, then settings.
    real_status = main.ENGINE.status
    main.ENGINE.status = lambda: {"model_downloaded": True, "selected_model": "large-v3-turbo",
                                  "state": "idle", "ready": False, "loaded_model": "",
                                  "target_device": "cpu", "loaded_device": ""}
    try:
        doc_id = projects.add_document(a, "x.wav", data=b"RIFF synthetic")["id"]
        url = f"/api/projects/{a}/transcribe"

        def queued_language(**body):
            projects.set_status(a, doc_id, "uploaded", "")
            r = client.post(url, json={"documents": [doc_id], **body})
            return main.QUEUE.get(r.json()["queued"][0]["job"]).payload["language"]

        client.put(f"/api/projects/{a}", json={"transcription_language": "es"})
        check("the project's language is used when none is sent", queued_language() == "es")
        check("an explicit request beats the project", queued_language(language="de") == "de")
        client.put(f"/api/projects/{a}", json={"transcription_language": ""})
        check("with no project choice the AI Settings default is used", queued_language() == "fr")
    finally:
        main.ENGINE.status = real_status

    SETTINGS.update({"task": "translate"})
    check("a global translate default does not change an existing project's operation",
          projects.project_operation(projects.load_project(a)) == "transcribe")
    SETTINGS.update({"task": "transcribe"})
    SETTINGS.update({"language": "auto"})


# ---------------------------------------------------------------------------

def main() -> int:
    tests = [
        test_parse_and_names,
        test_choices,
        test_codes_match_whisper_tokenizer,
        test_model_support,
        test_join_tokens,
        test_word_gaps_preserve_original_spacing,
        test_setting_schema_and_migration,
        test_engine_params,
        test_mlx_params,
        test_engine_reports_language,
        test_language_persists,
        test_legacy_documents_still_load,
        test_word_count_by_script,
        test_job_queue_carries_language,
        test_language_flows_through_the_job,
        test_unicode_exports,
        test_chinese_exports_and_captions,
        test_wrapping_spaceless_text,
        test_retimestamp_spaceless,
        test_translate_task_is_not_a_source_language_transcript,
        test_translate_task_survives_the_job_pipeline,
        test_right_to_left_languages,
        test_proofreading_is_english_only,
        test_api_and_pages,
        test_project_language_is_per_project,
    ]
    try:
        for test in tests:
            print(f"\n{test.__name__}")
            test()
    finally:
        shutil.rmtree(_TMP, ignore_errors=True)

    print()
    if FAILURES:
        print(f"{len(FAILURES)} check(s) failed: {FAILURES}")
        return 1
    print("all checks passed")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
